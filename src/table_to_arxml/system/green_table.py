"""绿表的读取、收集与合并。

绿表按列号 A..K 取值：B 为 app 自身端口，C 为对端模块，H 为数据元素，I 为对端组件（总线行为通道），
J 为对端端口（总线行为信号），K 为备注。R-Port、P-Port 两页结构相同，只是端口方向相反。
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from ..common.openpyxl_compat import copy_worksheet
from ..common.paths import expand_path, filesystem_path
from .values import norm

LOGGER = logging.getLogger(__name__)

R_PORT_SHEET = "R-Port"
P_PORT_SHEET = "P-Port"
LEGACY_INPUT_SHEET = "INPUT"
LEGACY_OUTPUT_SHEET = "OUTPUT"
_GREEN_SHEET_PAIRS = (
    (R_PORT_SHEET, P_PORT_SHEET),
    (LEGACY_INPUT_SHEET, LEGACY_OUTPUT_SHEET),
)
GREEN_COLUMN_COUNT = 11
# 表头只校验关键列，防止列错位；空串表示该列只要求非空。
GREEN_REQUIRED_HEADERS = {1: "", 2: "", 8: "component", 9: "port"}
RESERVED_REMARK_KEYWORDS = ("预留", "保留", "备用")


@dataclass(frozen=True)
class GreenRow:
    """绿表 R-Port/P-Port 页的一行有效数据。

    ``app_port`` 在 R-Port 页是 app 的 R-port，在 P-Port 页是 P-port；``peer``、``peer_port``
    是对端组件和端口，对端模块为 CAN/LIN 时分别是通道名和信号名。
    """

    app_port: str
    module: str
    peer: str
    peer_port: str
    data_element: str
    remark: str
    excel_row: int


@dataclass(frozen=True)
class GreenWorkbookInfo:
    """单个绿表抽象后的完整信息。"""

    app_name: str
    r_port_rows: list[GreenRow]
    p_port_rows: list[GreenRow]


def _green_sheet_names(sheetnames: list[str], green_path: Path) -> tuple[str, str]:
    """返回实际使用的 R-Port/P-Port 页签名；仅兼容旧 INPUT/OUTPUT 工作簿。"""

    present = set(sheetnames)
    available = [pair for pair in _GREEN_SHEET_PAIRS if set(pair) <= present]
    recognized = {name for pair in _GREEN_SHEET_PAIRS for name in pair} & present
    if len(available) != 1 or len(recognized) != 2:
        choices = f"{R_PORT_SHEET}/{P_PORT_SHEET}"
        legacy = f"{LEGACY_INPUT_SHEET}/{LEGACY_OUTPUT_SHEET}"
        raise ValueError(
            f"Green table must contain exactly one complete {choices} or legacy {legacy} sheet pair: "
            f"{green_path}"
        )
    return available[0]


def _is_reserved_remark(remark: str) -> bool:
    return any(keyword in remark for keyword in RESERVED_REMARK_KEYWORDS)


def _check_green_header(ws, green_path: Path, sheet: str) -> None:
    header = [norm(value) for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    header += [""] * (GREEN_COLUMN_COUNT - len(header))
    bad = [
        f"{get_column_letter(index + 1)}='{header[index]}'"
        for index, keyword in GREEN_REQUIRED_HEADERS.items()
        if not header[index] or keyword not in header[index].lower()
    ]
    if bad:
        raise ValueError(
            f"{green_path} {sheet}: invalid header {', '.join(bad)}; expected columns A..K with "
            "B=port, C=module, I=*Component, J=*port")


def _read_green_rows(ws, location: str, errors: list[str]) -> list[GreenRow]:
    rows: list[GreenRow] = []
    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        app_port, module, data_element, peer, peer_port, remark = (
            norm(row[index]) if index < len(row) else "" for index in (1, 2, 7, 8, 9, 10))
        if not any([app_port, module, peer, peer_port]):
            continue
        # I/J 同时为空时静默跳过；仅一列为空且备注未说明预留用途时报错。
        if not peer and not peer_port:
            continue
        if not peer or not peer_port:
            if not _is_reserved_remark(remark):
                errors.append(f"{location} row {excel_row}: column I or J is empty")
            continue
        if not app_port:
            errors.append(f"{location} row {excel_row}: column B is empty")
            continue
        rows.append(GreenRow(app_port, module, peer, peer_port, data_element, remark, excel_row))
    return rows


def parse_green_workbook(green_path: Path, app_name: str) -> GreenWorkbookInfo:
    """读取单个绿表的 R-Port/P-Port 页；兼容旧页名，逐行记录全部错误后统一失败。

    按固定列号取值而不依赖表头文案，避免换行、空格、全角字符差异导致匹配失败。
    """

    workbook = load_workbook(green_path, read_only=True, data_only=True)
    try:
        receive_sheet, send_sheet = _green_sheet_names(workbook.sheetnames, green_path)
        errors: list[str] = []
        rows: dict[str, list[GreenRow]] = {}
        for sheet in (receive_sheet, send_sheet):
            _check_green_header(workbook[sheet], green_path, sheet)
            rows[sheet] = _read_green_rows(workbook[sheet], f"{green_path} {sheet}", errors)
    finally:
        workbook.close()
    for message in errors:
        LOGGER.error("%s", message)
    if errors:
        raise ValueError(f"{green_path}: {len(errors)} invalid row(s)")
    return GreenWorkbookInfo(app_name, rows[receive_sheet], rows[send_sheet])


def collect_green_workbooks(
    entries: Iterable[str],
    variant: str,
    base_dir: Path,
) -> list[tuple[str, Path]]:
    """展开绿表路径和 glob，以文件名 stem 作为 app 名并去重。"""

    result: list[tuple[str, Path]] = []
    seen_paths: set[Path] = set()
    for entry in entries:
        rendered = norm(entry).replace("${variant}", variant)
        if not rendered:
            continue
        pattern = filesystem_path(rendered)
        matches = expand_path(pattern, base_dir)

        for match in sorted(matches, key=lambda path: str(path).casefold()):
            path = match.resolve()
            if path.name.startswith("~$"):
                continue
            if path.suffix.casefold() != ".xlsx" or not path.is_file():
                raise FileNotFoundError(f"Green table does not exist: {path}")
            if path not in seen_paths:
                seen_paths.add(path)
                result.append((path.stem, path))

    if not result:
        raise ValueError("No green tables were selected")
    return result


MAX_SHEET_NAME = 31


def _merged_sheet_name(app_name: str, sheet: str, used: set[str]) -> str:
    """生成合并工作簿中的页签名：``<app>_<页>``，超过 Excel 的 31 字符上限时截断；
    截断后与已有页签重名（不区分大小写）则追加 ``~序号``。"""

    base = f"{app_name}_{sheet}"
    candidate = base[:MAX_SHEET_NAME]
    index = 1
    while candidate.casefold() in used:
        index += 1
        suffix = f"~{index}"
        candidate = base[:MAX_SHEET_NAME - len(suffix)] + suffix
    used.add(candidate.casefold())
    return candidate


def merge_green_workbooks(
    green_workbooks: list[tuple[str, Path]],
    output_file: Path,
) -> Path:
    """把本次处理到的所有绿表合并成一个工作簿，便于统一检索。"""

    output_file.parent.mkdir(parents=True, exist_ok=True)

    merged_workbook = Workbook()
    placeholder_sheet = merged_workbook.active
    if placeholder_sheet is None:
        raise RuntimeError("Cannot access default sheet in merged workbook")
    placeholder_sheet.title = "MergedGreenTables"

    used_names: set[str] = set()
    for index, (app_name, green_workbook_path) in enumerate(green_workbooks):
        source_workbook = load_workbook(green_workbook_path, read_only=False, data_only=False)
        receive_sheet_name, send_sheet_name = _green_sheet_names(
            source_workbook.sheetnames, green_workbook_path,
        )
        receive_name = _merged_sheet_name(app_name, R_PORT_SHEET, used_names)
        if index == 0:
            receive_sheet = placeholder_sheet
            receive_sheet.title = receive_name
        else:
            receive_sheet = merged_workbook.create_sheet(receive_name)

        send_sheet = merged_workbook.create_sheet(_merged_sheet_name(app_name, P_PORT_SHEET, used_names))
        copy_worksheet(receive_sheet, source_workbook[receive_sheet_name])
        copy_worksheet(send_sheet, source_workbook[send_sheet_name])
        source_workbook.close()

    merged_workbook.save(output_file)
    merged_workbook.close()
    return output_file
