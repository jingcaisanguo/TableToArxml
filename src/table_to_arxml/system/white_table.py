"""白表（System Excel）的布局、创建、格式化与读取。

白表完全是本工具的产物，不依赖模板：页签与表头由 :data:`WHITE_TABLE_LAYOUT` 固定。
读取时只要求布局中的列各出现一次，额外的业务信息列不参与转换。
"""

from __future__ import annotations

from dataclasses import dataclass
from math import ceil
from pathlib import Path
from unicodedata import east_asian_width

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .values import norm

GENERAL_SHEET = "General"
COMPONENTS_SHEET = "Components"
ASSEMBLY_SHEET = "AssemblySwConnection"
DELEGATION_SHEET = "DelegationSwConnection"
SIGNAL_SHEET = "SenderReceiverToSignal"
SIGNAL_GROUP_SHEET = "SenderReceiverToSignalGroup"
WHITE_TABLE_LAYOUT = {
    GENERAL_SHEET: ["Key", "Value"],
    COMPONENTS_SHEET: ["Prototype", "SW-C", "Type", "Action"],
    ASSEMBLY_SHEET: ["Provider SW-C", "PPort", "Requester SW-C", "RPort"],
    DELEGATION_SHEET: ["Inner SW-C", "Inner Port", "Outer Port"],
    SIGNAL_SHEET: ["Signal", "Direction", "SW-C", "Port", "DataElement"],
    SIGNAL_GROUP_SHEET: ["SignalGroup", "SW-C", "Port", "DataElement"],
}
REQUIRED_SHEETS = (GENERAL_SHEET, COMPONENTS_SHEET)
GENERAL_KEYS = ("Composition", "System", "RootSwCompositionPrototype", "EcuInstance")


@dataclass(frozen=True)
class TableRow:
    sheet: str
    number: int
    values: dict[str, str]

    @property
    def location(self) -> str:
        return f"{self.sheet} row {self.number}"


@dataclass(frozen=True)
class CompositionInput:
    general: dict[str, str]
    tables: dict[str, list[TableRow]]


def read_composition_input(path: Path) -> CompositionInput:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        tables: dict[str, list[TableRow]] = {}
        for sheet_name, required_headers in WHITE_TABLE_LAYOUT.items():
            tables[sheet_name] = []
            if sheet_name not in workbook.sheetnames:
                if sheet_name in REQUIRED_SHEETS:
                    raise ValueError(f"Missing sheet '{sheet_name}' in {path}")
                continue
            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = [norm(value) for value in next(rows, ())]
            for header in required_headers:
                if headers.count(header) != 1:
                    raise ValueError(f"{path} {sheet_name}: expected exactly one '{header}' column")
            indices = {header: headers.index(header) for header in required_headers}
            for number, row in enumerate(rows, start=2):
                values = {
                    header: norm(row[index]) if index < len(row) else ""
                    for header, index in indices.items()
                }
                if any(values.values()):
                    tables[sheet_name].append(TableRow(sheet_name, number, values))

        general: dict[str, str] = {}
        for row in tables[GENERAL_SHEET]:
            key, value = row.values["Key"], row.values["Value"]
            if key in general and general[key] != value:
                raise ValueError(f"{row.location}: conflicting General key '{key}'")
            general[key] = value
        missing = [key for key in GENERAL_KEYS if not general.get(key)]
        if missing:
            raise ValueError(f"Missing General values: {', '.join(missing)}")
        return CompositionInput(general, tables)
    finally:
        workbook.close()


def create_white_workbook(general: dict[str, str]) -> Workbook:
    """按固定布局创建空白表，并用解析后的 General 值填充 General 页。"""

    workbook = Workbook()
    default_sheet = workbook.active
    if default_sheet is None:
        raise RuntimeError("Cannot access default sheet in white workbook")
    workbook.remove(default_sheet)

    for sheet_name, header in WHITE_TABLE_LAYOUT.items():
        workbook.create_sheet(sheet_name).append(header)

    general_ws = workbook[GENERAL_SHEET]
    for key in GENERAL_KEYS:
        value = norm(general.get(key))
        if value:
            general_ws.append([key, value])
    return workbook


def _display_width(text: str) -> float:
    return sum(2 if east_asian_width(character) in ("W", "F") else 1.15 for character in text)


def format_white_workbook(workbook: Workbook) -> None:
    """只调整白表外观，按内容设置列宽及长文本换行后的行高。"""

    header_fill = PatternFill("solid", fgColor="245A64")
    stripe_fill = PatternFill("solid", fgColor="F0F5F6")
    plain_fill = PatternFill("solid", fgColor="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=11, color="243238")
    grid_side = Side(style="thin", color="B8C9CE")
    separator = Border(left=grid_side, right=grid_side, top=grid_side, bottom=grid_side)
    port_headers = {"PPort", "RPort", "Port", "Inner Port", "Outer Port", "DataElement"}
    path_headers = {"Type", "Signal", "SignalGroup", "Value"}

    for sheet_name, headers in WHITE_TABLE_LAYOUT.items():
        sheet = workbook[sheet_name]
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = True
        sheet.sheet_properties.tabColor = "245A64"
        sheet.print_title_rows = "1:1"
        sheet.auto_filter.ref = sheet.dimensions if sheet_name != GENERAL_SHEET else None
        sheet.row_dimensions[1].height = 28

        widths: dict[int, float] = {}
        for column_index, header in enumerate(headers, start=1):
            longest = max(
                _display_width(line)
                for row in sheet.iter_rows(min_col=column_index, max_col=column_index)
                for line in (str(row[0].value) if row[0].value is not None else "").split("\n")
            )
            minimum = 32 if header in port_headers else 40 if header in path_headers else 20
            if header in ("Action", "Direction"):
                minimum = 14
            width = min(100, max(minimum, ceil(longest) + 4))
            widths[column_index] = width
            sheet.column_dimensions[get_column_letter(column_index)].width = width
            cell = sheet.cell(1, column_index)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = separator
            cell.alignment = Alignment(horizontal="left", vertical="center")

        for row in sheet.iter_rows(min_row=2):
            line_count = 1
            for cell in row:
                cell.font = body_font
                cell.fill = stripe_fill if cell.row % 2 == 0 else plain_fill
                cell.border = separator
                centered = headers[cell.column - 1] in ("Action", "Direction")
                cell.alignment = Alignment(
                    horizontal="center" if centered else "left", vertical="center", wrap_text=True)
                text = str(cell.value) if cell.value is not None else ""
                wrapped_lines = sum(
                    max(1, ceil(_display_width(line) / (widths[cell.column] - 4)))
                    for line in text.split("\n"))
                line_count = max(line_count, wrapped_lines)
            sheet.row_dimensions[row[0].row].height = min(409, max(22, line_count * 15 + 6))
