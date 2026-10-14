"""绿表 -> 白表（System Excel）转换。

白表完全是本工具的输出产物，不依赖任何模板文件：页签与表头由
:data:`WHITE_TABLE_LAYOUT` 固定，供本地白表读取器复用，General 页的值来自配置。

绿表中的端口/信号关系写入以下 3 个页签：
1. Components
2. AssemblySwConnection
3. SenderReceiverToSignal

处理原则：
- 写入前做去重，避免重复记录
- 普通组件连接写入 AssemblySwConnection
- CAN/LIN 直连写入 SenderReceiverToSignal
- E2E wrapper 只参与 Components/AssemblySwConnection
"""

from __future__ import annotations

import glob
import logging
import re
from copy import copy, deepcopy
from dataclasses import dataclass
from math import ceil
from pathlib import Path
from typing import Iterable
from unicodedata import east_asian_width

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .arxml_model import ArxmlModel
from .file_list import filesystem_path

INPUT_SHEET = "INPUT"
OUTPUT_SHEET = "OUTPUT"
BUS_MODULES = {"CAN", "LIN"}
E2E_WRAPPER_COMPONENT = "E2EWrapper"
MULTI_VALUE_SEPARATOR = "#"
RESERVED_REMARK_KEYWORDS = ("预留", "保留", "备用")
DEFAULT_PROTOTYPE_PREFIX = "its"
# 绿表按列号 A..K 取值；表头只校验关键列，防止列错位。
GREEN_COLUMN_COUNT = 11
GREEN_REQUIRED_HEADERS = {1: "", 2: "", 8: "component", 9: "port"}
WHITE_TABLE_SHEETS = {
    "components": "Components",
    "assembly": "AssemblySwConnection",
    "signal": "SenderReceiverToSignal",
}
WHITE_TABLE_LAYOUT = {
    "General": ["Key", "Value"],
    "Components": ["Prototype", "SW-C", "Type", "Action"],
    "AssemblySwConnection": ["Connection", "Provider SW-C", "PPort", "Requester SW-C", "RPort"],
    "DelegationSwConnection": ["Connection", "Inner SW-C", "Inner Port", "Outer Port"],
    "SenderReceiverToSignal": ["Signal", "Direction", "SW-C", "Port", "DataElement"],
    "SenderReceiverToSignalGroup": ["SignalGroup", "SW-C", "Port", "DataElement"],
}
GENERAL_KEYS = ("Composition", "System", "RootSwCompositionPrototype", "EcuInstance")


@dataclass(frozen=True)
class GreenInputRow:
    """绿表 INPUT 页的一行有效数据。

    字段对应关系：
    - app_rport: 绿表 B 列，app 的 R-port
    - source_module: 绿表 C 列，输入模块类型，例如 CAN/LIN/CFG/其他组件
    - source_component: 绿表 I 列，上游组件名；若 C 列为 CAN/LIN，则存放通道名
    - provider_pport: 绿表 J 列，上游 P-port；若 C 列为 CAN/LIN，则存放信号名
    - data_element: 绿表 H 列；转 SenderReceiverToSignal 时写入白表 DataElement
    - remark: 绿表 K 列备注
    - excel_row: Excel 实际行号，主要用于日志提示
    """

    app_rport: str
    source_module: str
    source_component: str
    provider_pport: str
    data_element: str
    remark: str
    excel_row: int


@dataclass(frozen=True)
class GreenOutputRow:
    """绿表 OUTPUT 页的一行有效数据，字段与 INPUT 页对称。"""

    app_pport: str
    target_module: str
    target_component: str
    requester_rport: str
    data_element: str
    remark: str
    excel_row: int


@dataclass(frozen=True)
class GreenWorkbookInfo:
    """单个绿表抽象后的完整信息。"""

    app_name: str
    input_rows: list[GreenInputRow]
    output_rows: list[GreenOutputRow]


def norm(value: object) -> str:
    """统一把 Excel / XML 读取到的值转换成可用字符串。"""

    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def load_channel_sets(model: ArxmlModel) -> dict[str, frozenset[str]]:
    """从已加载的 ARXML 中读取各总线类型的物理通道集合。"""

    channels = model.channels()
    channel_sets = {
        module_name: frozenset(
            name for name, module in channels.items() if module == module_name)
        for module_name in BUS_MODULES
    }
    return channel_sets


def _split_multi_value(raw: str) -> list[str]:
    """拆分绿表多值字段。"""

    return [item.strip() for item in raw.split(MULTI_VALUE_SEPARATOR) if item and item.strip()]


def _is_reserved_remark(remark: str) -> bool:
    """判断备注是否包含保留关键字。"""

    return any(keyword in remark for keyword in RESERVED_REMARK_KEYWORDS)


def _build_signal_path(
    channel: str, signal: str, template: str = "/${channel}/Signal/${signal}",
) -> str:
    """把通道名和信号名拼成白表 SenderReceiverToSignal 所需的全路径。"""

    if signal.startswith("/"):
        return signal
    return template.replace("${channel}", channel).replace("${signal}", signal)


class ComponentTypeMap(dict):
    def __init__(self, candidates: dict[str, set[tuple[str, str]]]) -> None:
        super().__init__((alias, next(iter(values))) for alias, values in candidates.items()
                         if len(values) == 1)
        self.ambiguous = {alias: values for alias, values in candidates.items() if len(values) > 1}

    def get(self, key, default=None):
        if key in self.ambiguous:
            choices = ", ".join(f"{name} ({path})" for name, path in sorted(self.ambiguous[key]))
            raise ValueError(f"Ambiguous component '{key}': {choices}; configure componentAliases")
        return super().get(key, default)


def build_component_type_map(
    model: ArxmlModel,
    composition_path: str,
    aliases: dict[str, str | dict[str, str]] | None = None,
    prototype_prefix: str = DEFAULT_PROTOTYPE_PREFIX,
) -> ComponentTypeMap:
    """按完整路径索引类型，短名只在唯一时可用；已有实例限定到目标 composition。"""

    component_types = model.component_types()
    # 组合类型只能通过 componentAliases 显式引用，避免目标/根组合参与短名匹配。
    composition_types = model.composition_types()
    prototypes = model.component_prototypes(composition_path.rstrip("/"))
    candidates: dict[str, set[tuple[str, str]]] = {}
    by_type: dict[str, set[tuple[str, str]]] = {}
    named_types: set[str] = set()
    for instance_path, type_path in prototypes.items():
        definition = (instance_path.rsplit("/", 1)[-1], type_path)
        by_type.setdefault(type_path, set()).add(definition)
        candidates.setdefault(instance_path, set()).add(definition)
    for type_path, short_name in component_types.items():
        definitions = by_type.setdefault(type_path, {(prototype_prefix + short_name, type_path)})
        candidates.setdefault(type_path, set()).update(definitions)
        candidates.setdefault(short_name, set()).update(definitions)
        for definition in definitions:
            candidates.setdefault(definition[0], set()).add(definition)
    for alias, target in (aliases or {}).items():
        if not isinstance(alias, str) or not alias or "/" in alias or "\\" in alias:
            raise ValueError("componentAliases keys must be nonempty names without path separators")
        # 对象形式 {"type": 类型路径, "prototype": 实例名} 用于不带 prototypePrefix 的新实例名。
        prototype_name = ""
        path = target
        if isinstance(target, dict):
            path, prototype_name = target.get("type"), target.get("prototype")
            if not isinstance(prototype_name, str) or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", prototype_name):
                raise ValueError(f"componentAliases.{alias}.prototype must be an AUTOSAR SHORT-NAME")
        if not isinstance(path, str) or not path.startswith("/"):
            raise ValueError(f"componentAliases.{alias} must be a full AUTOSAR path")
        if path in prototypes and not prototype_name:
            definition = (path.rsplit("/", 1)[-1], prototypes[path])
        elif path in component_types or path in composition_types:
            definitions = by_type.setdefault(path, {(prototype_prefix + path.rsplit("/", 1)[-1], path)})
            if len(definitions) != 1:
                raise ValueError(f"componentAliases.{alias}: multiple target instances; use a prototype path")
            definition = next(iter(definitions))
            if path not in prototypes.values() and path not in named_types:
                definition = (prototype_name or prototype_prefix + alias, path)
                old_definition = next(iter(definitions))
                for values in candidates.values():
                    if old_definition in values:
                        values.remove(old_definition)
                        values.add(definition)
                by_type[path] = {definition}
                named_types.add(path)
            elif prototype_name and definition[0] != prototype_name:
                raise ValueError(
                    f"componentAliases.{alias}: {path} already has instance '{definition[0]}'")
        else:
            raise ValueError(f"componentAliases.{alias}: unknown type or prototype outside target: {path}")
        candidates[alias] = {definition}
        candidates.setdefault(definition[0], set()).add(definition)
    return ComponentTypeMap(candidates)


def parse_green_workbook(green_path: Path, app_name: str) -> GreenWorkbookInfo:
    """读取单个绿表，抽取 INPUT/OUTPUT 两页的关键字段。

    这里不依赖表头文案，而是按固定列号取值，避免表头换行、空格、全角字符差异
    导致匹配失败。
    """

    workbook = load_workbook(green_path, read_only=True, data_only=True)
    try:
        return _parse_green_workbook(workbook, green_path, app_name)
    finally:
        workbook.close()


def _check_green_header(ws, green_path: Path, sheet: str) -> None:
    header = [norm(value) for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    header += [""] * (GREEN_COLUMN_COUNT - len(header))
    bad = [
        f"{get_column_letter(index + 1)}='{header[index]}'"
        for index, keyword in GREEN_REQUIRED_HEADERS.items()
        if not header[index] or keyword not in header[index].lower()
    ]
    if bad:
        raise ValueError(
            f"{green_path} {sheet}: invalid header {', '.join(bad)}; expected columns A..K with "
            "B=port, C=module, I=*Component, J=*port")


def _parse_green_workbook(workbook, green_path: Path, app_name: str) -> GreenWorkbookInfo:

    if INPUT_SHEET not in workbook.sheetnames or OUTPUT_SHEET not in workbook.sheetnames:
        raise ValueError(
            f"Green table must contain {INPUT_SHEET}/{OUTPUT_SHEET} sheets: {green_path}")
    for sheet in (INPUT_SHEET, OUTPUT_SHEET):
        _check_green_header(workbook[sheet], green_path, sheet)

    input_rows: list[GreenInputRow] = []
    for excel_row, row in enumerate(
            workbook[INPUT_SHEET].iter_rows(min_row=2, values_only=True), start=2):
        app_rport = norm(row[1] if len(row) > 1 else "")
        source_module = norm(row[2] if len(row) > 2 else "")
        data_element = norm(row[7] if len(row) > 7 else "")
        source_component = norm(row[8] if len(row) > 8 else "")
        provider_pport = norm(row[9] if len(row) > 9 else "")
        remark = norm(row[10] if len(row) > 10 else "")

        if not any([app_rport, source_module, source_component, provider_pport]):
            continue

        # I/J 同时为空时静默跳过；仅一列为空且备注未说明预留用途时记录错误。
        if not source_component and not provider_pport:
            continue
        if not source_component or not provider_pport:
            if not _is_reserved_remark(remark):
                raise ValueError(f"{green_path} INPUT row {excel_row}: column I or J is empty")
            continue
        if not app_rport:
            raise ValueError(f"{green_path} INPUT row {excel_row}: column B is empty")

        input_rows.append(
            GreenInputRow(
                app_rport=app_rport,
                source_module=source_module,
                source_component=source_component,
                provider_pport=provider_pport,
                data_element=data_element,
                remark=remark,
                excel_row=excel_row,
            )
        )

    output_rows: list[GreenOutputRow] = []
    for excel_row, row in enumerate(
            workbook[OUTPUT_SHEET].iter_rows(min_row=2, values_only=True), start=2):
        app_pport = norm(row[1] if len(row) > 1 else "")
        target_module = norm(row[2] if len(row) > 2 else "")
        data_element = norm(row[7] if len(row) > 7 else "")
        target_component = norm(row[8] if len(row) > 8 else "")
        requester_rport = norm(row[9] if len(row) > 9 else "")
        remark = norm(row[10] if len(row) > 10 else "")

        if not any([app_pport, target_module, target_component, requester_rport]):
            continue

        if not target_component and not requester_rport:
            continue
        if not target_component or not requester_rport:
            if not _is_reserved_remark(remark):
                raise ValueError(f"{green_path} OUTPUT row {excel_row}: column I or J is empty")
            continue
        if not app_pport:
            raise ValueError(f"{green_path} OUTPUT row {excel_row}: column B is empty")

        output_rows.append(
            GreenOutputRow(
                app_pport=app_pport,
                target_module=target_module,
                target_component=target_component,
                requester_rport=requester_rport,
                data_element=data_element,
                remark=remark,
                excel_row=excel_row,
            )
        )

    return GreenWorkbookInfo(app_name=app_name, input_rows=input_rows, output_rows=output_rows)


def _load_existing_component_rows(ws) -> set[tuple[str, str, str, str]]:
    """读取白表 Components 页中已存在的完整记录键。"""

    return {
        (norm(row[0]), norm(row[1]), norm(row[2]), norm(row[3]))
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row and any(norm(cell) for cell in row[:4])
    }


def _load_existing_assembly_rows(ws) -> set[tuple[str, str, str, str]]:
    """读取白表 AssemblySwConnection 页中已有连线。"""

    return {
        (norm(row[1]), norm(row[2]), norm(row[3]), norm(row[4]))
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row and len(row) > 4 and any(norm(cell) for cell in row[1:5])
    }


def _load_existing_signal_rows(ws) -> set[tuple[str, str, str, str, str]]:
    """读取白表 SenderReceiverToSignal 页中已有信号映射。"""

    return {
        (norm(row[0]), norm(row[1]), norm(row[2]), norm(row[3]), norm(row[4]))
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row and len(row) > 4 and any(norm(cell) for cell in row[:5])
    }


def _append_component(
    ws,
    existing_rows: set[tuple[str, str, str, str]],
    prototype: str,
    component_name: str,
    component_type: str,
    action: str,
) -> bool:
    """向 Components 页追加一行组件定义，返回是否实际写入。"""

    row_key = (prototype, component_name, component_type, action)
    if row_key in existing_rows:
        return False
    if any(row[0] == prototype for row in existing_rows):
        raise ValueError(f"Conflicting component prototype '{prototype}'")
    ws.append([prototype, component_name, component_type, action])
    existing_rows.add(row_key)
    return True


def _append_assembly(
    ws,
    existing_rows: set[tuple[str, str, str, str]],
    provider: str,
    pport: str,
    requester: str,
    rport: str,
) -> bool:
    """向 AssemblySwConnection 页追加一条组件端口连线。"""

    row_key = (provider, pport, requester, rport)
    if not all(row_key) or row_key in existing_rows:
        return False
    ws.append([None, provider, pport, requester, rport])
    existing_rows.add(row_key)
    return True


def _append_signal_mapping(
    ws,
    existing_rows: set[tuple[str, str, str, str, str]],
    signal_path: str,
    direction: str,
    swc_name: str,
    port_name: str,
    data_element: str,
) -> bool:
    """向 SenderReceiverToSignal 页追加一条信号映射。"""

    row_key = (signal_path, direction, swc_name, port_name, data_element)
    if not all(row_key) or row_key in existing_rows:
        return False
    ws.append(list(row_key))
    existing_rows.add(row_key)
    return True


def _collect_non_app_components(
    info: GreenWorkbookInfo, e2e_wrapper: str = E2E_WRAPPER_COMPONENT,
) -> set[str]:
    """收集绿表中除 app 自身外，需要出现在白表 Components 页的组件名。

    CAN/LIN 行本身不是组件连线，不把通道名当普通组件处理；
    配置的 E2E wrapper 例外，按普通组件处理。
    """

    names: set[str] = set()
    for row in info.input_rows:
        module_name = row.source_module.upper()
        if module_name in BUS_MODULES and row.source_component != e2e_wrapper:
            continue
        if row.source_component:
            names.add(row.source_component)

    for row in info.output_rows:
        module_name = row.target_module.upper()
        if module_name in BUS_MODULES and row.target_component != e2e_wrapper:
            continue
        if row.target_component:
            names.add(row.target_component)
    return names


def process_green_table(
    info: GreenWorkbookInfo,
    component_type_map: dict[str, tuple[str, str]],
    app_names: set[str],
    channel_sets: dict[str, frozenset[str]],
    workbook,
    e2e_wrapper: str = E2E_WRAPPER_COMPONENT,
    signal_path_template: str = "/${channel}/Signal/${signal}",
    model: ArxmlModel | None = None,
    ecu_path: str = "",
) -> dict[str, int]:
    """把一个绿表对象写入白表工作簿。

    核心处理顺序：
    1. 先补 Components：绿表定义的 app 写成 Add，其他被引用组件写成 Reference
    2. 再处理 INPUT 页：普通组件 -> AssemblySwConnection，CAN/LIN -> SenderReceiverToSignal
    3. 再处理 OUTPUT 页，规则与 INPUT 对称
    """

    components_ws = workbook[WHITE_TABLE_SHEETS["components"]]
    assembly_ws = workbook[WHITE_TABLE_SHEETS["assembly"]]
    signal_ws = workbook[WHITE_TABLE_SHEETS["signal"]]

    existing_component_rows = _load_existing_component_rows(components_ws)
    existing_assembly_rows = _load_existing_assembly_rows(assembly_ws)
    existing_signal_rows = _load_existing_signal_rows(signal_ws)

    counters = {"components": 0, "assembly": 0, "signal": 0}
    app_component_info = component_type_map.get(info.app_name)
    if not app_component_info:
        raise ValueError(
            f"Cannot find the SW component type of green table '{info.app_name}' in the "
            "loaded ARXML files")
    app_prototype, app_component_type = app_component_info
    app_component_name = app_prototype
    app_definitions = {component_type_map.get(name) for name in app_names}

    if _append_component(
        components_ws,
        existing_component_rows,
        app_prototype,
        app_component_name,
        app_component_type,
        "Add",
    ):
        counters["components"] += 1

    for component_name in sorted(_collect_non_app_components(info, e2e_wrapper)):
        component_info = component_type_map.get(component_name)
        if not component_info:
            raise ValueError(f"Green table '{info.app_name}': unresolved component '{component_name}'")
        component_prototype, component_type = component_info
        action = "Add" if component_info in app_definitions else "Reference"
        if _append_component(
            components_ws,
            existing_component_rows,
            component_prototype,
            component_prototype,
            component_type,
            action,
        ):
            counters["components"] += 1

    for row in info.input_rows:
        module_name = row.source_module.upper()
        # INPUT 页中，CAN/LIN 行是 app R-port 到总线信号的映射，不是组件装配关系。
        if module_name in BUS_MODULES and row.source_component != e2e_wrapper:
            channels = _split_multi_value(row.source_component)
            if not channels or not row.provider_pport:
                raise ValueError(f"{info.app_name} INPUT row {row.excel_row}: channel or signal is empty")
            for channel in channels:
                if channel not in channel_sets[module_name]:
                    raise ValueError(
                        f"{info.app_name} INPUT row {row.excel_row}: channel '{channel}' "
                        f"is not defined in the {module_name} ARXML")
                # 输入信号的白表 Direction 固定为 IN，DataElement 优先取绿表 H 列。
                signal_path = _build_signal_path(channel, row.provider_pport, signal_path_template)
                if model is None:
                    raise ValueError("ARXML model is required for CAN/LIN signal validation")
                try:
                    model.validate_signal_mapping(signal_path, "IN", ecu_path, module_name, channel)
                except ValueError as exc:
                    raise ValueError(f"{info.app_name} INPUT row {row.excel_row}: {exc}") from exc
                if _append_signal_mapping(
                    signal_ws,
                    existing_signal_rows,
                    signal_path,
                    "IN",
                    app_component_name,
                    row.app_rport,
                    row.data_element or row.app_rport,
                ):
                    counters["signal"] += 1
            continue

        # 非 CAN/LIN 场景：按“上游 provider -> 当前 app requester”写组件连线。
        component_name = component_type_map.get(row.source_component)[0]
        if _append_assembly(
            assembly_ws,
            existing_assembly_rows,
            component_name,
            row.provider_pport,
            app_component_name,
            row.app_rport,
        ):
            counters["assembly"] += 1

    for row in info.output_rows:
        module_name = row.target_module.upper()
        # OUTPUT 页中，CAN/LIN 行表示 app P-port 发往总线信号。
        if module_name in BUS_MODULES and row.target_component != e2e_wrapper:
            channels = _split_multi_value(row.target_component)
            if not channels or not row.requester_rport:
                raise ValueError(f"{info.app_name} OUTPUT row {row.excel_row}: channel or signal is empty")
            for channel in channels:
                if channel not in channel_sets[module_name]:
                    raise ValueError(
                        f"{info.app_name} OUTPUT row {row.excel_row}: channel '{channel}' "
                        f"is not defined in the {module_name} ARXML")
                # 输出信号的白表 Direction 固定为 OUT，DataElement 优先取绿表 H 列。
                signal_path = _build_signal_path(channel, row.requester_rport, signal_path_template)
                if model is None:
                    raise ValueError("ARXML model is required for CAN/LIN signal validation")
                try:
                    model.validate_signal_mapping(signal_path, "OUT", ecu_path, module_name, channel)
                except ValueError as exc:
                    raise ValueError(f"{info.app_name} OUTPUT row {row.excel_row}: {exc}") from exc
                if _append_signal_mapping(
                    signal_ws,
                    existing_signal_rows,
                    signal_path,
                    "OUT",
                    app_component_name,
                    row.app_pport,
                    row.data_element or row.app_pport,
                ):
                    counters["signal"] += 1
            continue

        # 非 CAN/LIN 场景：按“当前 app provider -> 下游 requester”写组件连线。
        component_name = component_type_map.get(row.target_component)[0]
        if _append_assembly(
            assembly_ws,
            existing_assembly_rows,
            app_component_name,
            row.app_pport,
            component_name,
            row.requester_rport,
        ):
            counters["assembly"] += 1

    return counters


def _build_merged_sheet_name(app_name: str, direction: str) -> str:
    """生成合并绿表中的页签名。"""

    sheet_name = f"{app_name}_{direction}"
    return sheet_name[:31]


def _copy_green_sheet(target_sheet, source_sheet) -> None:
    """复制单个绿表页签的内容、单元格样式和工作表格式。"""

    for (row_index, column_index), source_cell in source_sheet._cells.items():
        target_cell = target_sheet.cell(row=row_index, column=column_index)
        target_cell._value = source_cell._value
        target_cell.data_type = source_cell.data_type
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)
        target_cell.quotePrefix = source_cell.quotePrefix
        target_cell.pivotButton = source_cell.pivotButton
        if source_cell.hyperlink:
            target_cell._hyperlink = copy(source_cell.hyperlink)
        if source_cell.comment:
            target_cell.comment = copy(source_cell.comment)

    for key, dimension in source_sheet.row_dimensions.items():
        target_sheet.row_dimensions[key] = copy(dimension)
        target_sheet.row_dimensions[key].worksheet = target_sheet
    for key, dimension in source_sheet.column_dimensions.items():
        target_sheet.column_dimensions[key] = copy(dimension)
        target_sheet.column_dimensions[key].worksheet = target_sheet

    for merged_range in source_sheet.merged_cells.ranges:
        target_sheet.merge_cells(str(merged_range))

    target_sheet.sheet_format = copy(source_sheet.sheet_format)
    target_sheet.sheet_properties = copy(source_sheet.sheet_properties)
    target_sheet.page_margins = copy(source_sheet.page_margins)
    target_sheet.page_setup = copy(source_sheet.page_setup)
    target_sheet.print_options = copy(source_sheet.print_options)
    target_sheet.protection = copy(source_sheet.protection)
    target_sheet.views = deepcopy(source_sheet.views)
    target_sheet.freeze_panes = source_sheet.freeze_panes
    target_sheet.auto_filter = deepcopy(source_sheet.auto_filter)
    target_sheet.data_validations = deepcopy(source_sheet.data_validations)
    target_sheet.conditional_formatting = deepcopy(source_sheet.conditional_formatting)
    target_sheet.oddHeader = copy(source_sheet.oddHeader)
    target_sheet.oddFooter = copy(source_sheet.oddFooter)
    target_sheet.evenHeader = copy(source_sheet.evenHeader)
    target_sheet.evenFooter = copy(source_sheet.evenFooter)
    target_sheet.firstHeader = copy(source_sheet.firstHeader)
    target_sheet.firstFooter = copy(source_sheet.firstFooter)


def merge_green_workbooks(
    green_workbooks: list[tuple[str, Path]],
    output_file: Path,
) -> Path:
    """把本次处理到的所有绿表合并成一个工作簿，便于统一检索。"""

    output_file.parent.mkdir(parents=True, exist_ok=True)

    merged_workbook = Workbook()
    placeholder_sheet = merged_workbook.active
    if placeholder_sheet is None:
        raise RuntimeError("Cannot access default sheet in merged workbook")
    placeholder_sheet.title = "MergedGreenTables"

    for index, (app_name, green_workbook_path) in enumerate(green_workbooks):
        source_workbook = load_workbook(green_workbook_path, read_only=False, data_only=False)
        input_sheet_name = _build_merged_sheet_name(app_name, INPUT_SHEET)
        output_sheet_name = _build_merged_sheet_name(app_name, OUTPUT_SHEET)

        if index == 0:
            input_sheet = placeholder_sheet
            input_sheet.title = input_sheet_name
        else:
            input_sheet = merged_workbook.create_sheet(input_sheet_name)

        output_sheet = merged_workbook.create_sheet(output_sheet_name)
        _copy_green_sheet(input_sheet, source_workbook[INPUT_SHEET])
        _copy_green_sheet(output_sheet, source_workbook[OUTPUT_SHEET])
        source_workbook.close()

    merged_workbook.save(output_file)
    merged_workbook.close()
    return output_file


def resolve_general(general: dict, model: ArxmlModel) -> dict[str, str]:
    """解析白表 General 页的值。

    ``Composition`` 决定新组件放在哪里，必须配置；其余项未配置时按唯一性推导：
    ``System`` 取根组合逐层包含该 Composition 的 System，``EcuInstance`` 取该 System
    FIBEX-ELEMENTS 中的 ECU-INSTANCE，``RootSwCompositionPrototype`` 取该 System 的根组合。
    """

    if not isinstance(general, dict):
        raise ValueError("config.general must be an object")

    composition = norm(general.get("Composition"))
    if not composition:
        raise KeyError("Missing config.general.Composition")
    resolved: dict[str, str] = {"Composition": composition}

    system_name = norm(general.get("System"))
    if system_name:
        system = model.find_system(system_name)
    else:
        systems = model.systems_containing(composition)
        if len(systems) != 1:
            raise KeyError(
                f"Cannot derive config.general.System: {len(systems)} System(s) contain {composition}"
                + "".join(f", {item.getFullName()}" for item in systems))
        system = systems[0]
        logging.info("config.general.System resolved from Composition %s: %s", composition, system.getFullName())
    # 已有 System 统一写完整路径，生成时按同路径拆分合并到该 System。
    resolved["System"] = system.getFullName() if system is not None else system_name

    ecu_instance = norm(general.get("EcuInstance"))
    if not ecu_instance:
        ecu_instances = model.system_ecu_instances(system) if system is not None else []
        if len(ecu_instances) != 1:
            raise KeyError(
                f"Cannot derive config.general.EcuInstance: System '{resolved['System']}' "
                f"references {len(ecu_instances)} ECU-INSTANCE(s)")
        ecu_instance = ecu_instances[0]
        logging.info("config.general.EcuInstance resolved from System %s: %s", resolved["System"], ecu_instance)
    resolved["EcuInstance"] = ecu_instance

    root_composition = norm(general.get("RootSwCompositionPrototype"))
    if not root_composition:
        root_composition = model.root_composition_path(resolved["System"])
        if not root_composition:
            raise KeyError(
                "Missing config.general.RootSwCompositionPrototype because System "
                f"'{resolved['System']}' has no ROOT-SOFTWARE-COMPOSITION in the loaded ARXML")
        logging.info(
            "config.general.RootSwCompositionPrototype resolved from System %s: %s",
            resolved["System"], root_composition)
    resolved["RootSwCompositionPrototype"] = root_composition
    return resolved


def create_white_workbook(general: dict[str, str]) -> Workbook:
    """按固定布局创建空白表，并用解析后的 General 值填充 General 页。"""

    workbook = Workbook()
    default_sheet = workbook.active
    if default_sheet is None:
        raise RuntimeError("Cannot access default sheet in white workbook")
    workbook.remove(default_sheet)

    for sheet_name, header in WHITE_TABLE_LAYOUT.items():
        workbook.create_sheet(sheet_name).append(header)

    general_ws = workbook["General"]
    for key in GENERAL_KEYS:
        value = norm(general.get(key))
        if value:
            general_ws.append([key, value])
    return workbook


def _display_width(text: str) -> float:
    return sum(2 if east_asian_width(character) in ("W", "F") else 1.15 for character in text)


def format_white_workbook(workbook: Workbook) -> None:
    """只调整白表外观，按内容设置列宽及长文本换行后的行高。"""

    header_fill = PatternFill("solid", fgColor="245A64")
    stripe_fill = PatternFill("solid", fgColor="F0F5F6")
    plain_fill = PatternFill("solid", fgColor="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=11, color="243238")
    grid_side = Side(style="thin", color="B8C9CE")
    separator = Border(left=grid_side, right=grid_side, top=grid_side, bottom=grid_side)
    port_headers = {"PPort", "RPort", "Port", "Inner Port", "Outer Port", "DataElement"}
    path_headers = {"Type", "Signal", "SignalGroup", "Value"}

    for sheet_name, headers in WHITE_TABLE_LAYOUT.items():
        sheet = workbook[sheet_name]
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = True
        sheet.sheet_properties.tabColor = "245A64"
        sheet.print_title_rows = "1:1"
        sheet.auto_filter.ref = sheet.dimensions if sheet_name != "General" else None
        sheet.row_dimensions[1].height = 28

        widths: dict[int, float] = {}
        for column_index, header in enumerate(headers, start=1):
            longest = max(
                _display_width(line)
                for row in sheet.iter_rows(min_col=column_index, max_col=column_index)
                for line in (str(row[0].value) if row[0].value is not None else "").split("\n")
            )
            minimum = 32 if header in port_headers else 40 if header in path_headers else 20
            if header in ("Action", "Direction"):
                minimum = 14
            width = min(100, max(minimum, ceil(longest) + 4))
            widths[column_index] = width
            sheet.column_dimensions[get_column_letter(column_index)].width = width
            cell = sheet.cell(1, column_index)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = separator
            cell.alignment = Alignment(horizontal="left", vertical="center")

        for row in sheet.iter_rows(min_row=2):
            line_count = 1
            for cell in row:
                cell.font = body_font
                cell.fill = stripe_fill if cell.row % 2 == 0 else plain_fill
                cell.border = separator
                centered = headers[cell.column - 1] in ("Action", "Direction")
                cell.alignment = Alignment(
                    horizontal="center" if centered else "left", vertical="center", wrap_text=True)
                text = str(cell.value) if cell.value is not None else ""
                wrapped_lines = sum(
                    max(1, ceil(_display_width(line) / (widths[cell.column] - 4)))
                    for line in text.split("\n"))
                line_count = max(line_count, wrapped_lines)
            sheet.row_dimensions[row[0].row].height = min(409, max(22, line_count * 15 + 6))


def collect_green_workbooks(
    entries: Iterable[str],
    variant: str,
    base_dir: Path,
) -> list[tuple[str, Path]]:
    """展开绿表路径和 glob，以文件名 stem 作为 app 名并去重。"""

    result: list[tuple[str, Path]] = []
    seen_paths: set[Path] = set()
    for entry in entries:
        rendered = norm(entry).replace("${variant}", variant)
        if not rendered:
            continue
        pattern = filesystem_path(rendered)
        if not pattern.is_absolute():
            pattern = base_dir / pattern

        if glob.has_magic(str(pattern)):
            # 通配符未命中视为没有可选绿表，不当作配置错误。
            matches = [Path(match) for match in glob.glob(str(pattern), recursive=True)]
        else:
            matches = [pattern]

        for match in sorted(matches, key=lambda path: str(path).casefold()):
            path = match.resolve()
            if path.name.startswith("~$"):
                continue
            if path.suffix.casefold() != ".xlsx" or not path.is_file():
                raise FileNotFoundError(f"Green table does not exist: {path}")
            if path not in seen_paths:
                seen_paths.add(path)
                result.append((path.stem, path))

    if not result:
        raise ValueError("No green tables were selected")
    return result


def generate_white_table(
    model: ArxmlModel,
    green_tables: list[tuple[str, Path]],
    general: dict,
    output_file: Path,
    component_aliases: dict[str, str | dict[str, str]] | None = None,
    prototype_prefix: str = DEFAULT_PROTOTYPE_PREFIX,
    e2e_wrapper: str = E2E_WRAPPER_COMPONENT,
    signal_path_template: str = "/${channel}/Signal/${signal}",
) -> Path:
    """基于已加载的 ARXML 模型和绿表生成白表。"""

    channel_sets = load_channel_sets(model)
    component_type_map = build_component_type_map(
        model, norm(general.get("Composition")), component_aliases, prototype_prefix)
    general = resolve_general(general, model)
    workbook = create_white_workbook(general)

    # 一个绿表对应一个要创建的模块，模块类型必须能在 ARXML 中找到。
    missing_components = sorted(
        app_name for app_name, _ in green_tables if component_type_map.get(app_name) is None)
    if missing_components:
        raise ValueError(
            "Cannot find the SW component type of green table(s) in the loaded ARXML files: "
            + ", ".join(missing_components))

    total_counters = {"components": 0, "assembly": 0, "signal": 0}
    app_names = {app_name for app_name, _ in green_tables}
    for app_name, green_excel in green_tables:
        green_info = parse_green_workbook(green_excel, app_name)
        counters = process_green_table(
            green_info, component_type_map, app_names, channel_sets, workbook,
            e2e_wrapper, signal_path_template, model, norm(general.get("EcuInstance")))
        logging.info(
            "Processed %s: +%s Components, +%s AssemblySwConnection, +%s SenderReceiverToSignal",
            app_name,
            counters["components"],
            counters["assembly"],
            counters["signal"],
        )
        for key in total_counters:
            total_counters[key] += counters[key]

    format_white_workbook(workbook)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    workbook.close()
    logging.info(
        "White table done: %s Components, %s AssemblySwConnection, %s SenderReceiverToSignal -> %s",
        total_counters["components"],
        total_counters["assembly"],
        total_counters["signal"],
        output_file,
    )
    return output_file
