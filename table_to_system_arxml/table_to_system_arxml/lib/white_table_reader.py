"""Read the tool's white-table format without an external converter."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .green_to_white import GENERAL_KEYS, WHITE_TABLE_LAYOUT, norm


@dataclass(frozen=True)
class TableRow:
    sheet: str
    number: int
    values: dict[str, str]

    @property
    def location(self) -> str:
        return f"{self.sheet} row {self.number}"


@dataclass(frozen=True)
class CompositionInput:
    general: dict[str, str]
    tables: dict[str, list[TableRow]]


def read_composition_input(path: Path) -> CompositionInput:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        tables: dict[str, list[TableRow]] = {}
        for sheet_name, required_headers in WHITE_TABLE_LAYOUT.items():
            tables[sheet_name] = []
            if sheet_name not in workbook.sheetnames:
                if sheet_name in ("General", "Components"):
                    raise ValueError(f"Missing sheet '{sheet_name}' in {path}")
                continue
            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = [norm(value) for value in next(rows, ())]
            for header in required_headers:
                if headers.count(header) != 1:
                    raise ValueError(f"{path} {sheet_name}: expected exactly one '{header}' column")
            indices = {header: headers.index(header) for header in required_headers}
            for number, row in enumerate(rows, start=2):
                values = {
                    header: norm(row[index]) if index < len(row) else ""
                    for header, index in indices.items()
                }
                if any(values.values()):
                    tables[sheet_name].append(TableRow(sheet_name, number, values))

        general: dict[str, str] = {}
        for row in tables["General"]:
            key, value = row.values["Key"], row.values["Value"]
            if key in general and general[key] != value:
                raise ValueError(f"{row.location}: conflicting General key '{key}'")
            general[key] = value
        missing = [key for key in GENERAL_KEYS if not general.get(key)]
        if missing:
            raise ValueError(f"Missing General values: {', '.join(missing)}")
        return CompositionInput(general, tables)
    finally:
        workbook.close()