"""绿表的读取、收集与合并。

绿表按列号 A..K 取值：B 为 app 自身端口，C 为对端模块，H 为数据元素，I 为对端组件（总线行为通道），
J 为对端端口（总线行为信号），K 为备注。INPUT、OUTPUT 两页结构相同，只是端口方向相反。
"""

from __future__ import annotations

import glob
import logging
from copy import copy, deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .common import norm
from .file_list import filesystem_path

LOGGER = logging.getLogger(__name__)

INPUT_SHEET = "INPUT"
OUTPUT_SHEET = "OUTPUT"
GREEN_COLUMN_COUNT = 11
# 表头只校验关键列，防止列错位；空串表示该列只要求非空。
GREEN_REQUIRED_HEADERS = {1: "", 2: "", 8: "component", 9: "port"}
RESERVED_REMARK_KEYWORDS = ("预留", "保留", "备用")


@dataclass(frozen=True)
class GreenRow:
    """绿表 INPUT/OUTPUT 页的一行有效数据。

    ``app_port`` 在 INPUT 页是 app 的 R-port，在 OUTPUT 页是 P-port；``peer``、``peer_port``
    是对端组件和端口，对端模块为 CAN/LIN 时分别是通道名和信号名。
    """

    app_port: str
    module: str
    peer: str
    peer_port: str
    data_element: str
    remark: str
    excel_row: int


@dataclass(frozen=True)
class GreenWorkbookInfo:
    """单个绿表抽象后的完整信息。"""

    app_name: str
    input_rows: list[GreenRow]
    output_rows: list[GreenRow]


def _is_reserved_remark(remark: str) -> bool:
    return any(keyword in remark for keyword in RESERVED_REMARK_KEYWORDS)


def _check_green_header(ws, green_path: Path, sheet: str) -> None:
    header = [norm(value) for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    header += [""] * (GREEN_COLUMN_COUNT - len(header))
    bad = [
        f"{get_column_letter(index + 1)}='{header[index]}'"
        for index, keyword in GREEN_REQUIRED_HEADERS.items()
        if not header[index] or keyword not in header[index].lower()
    ]
    if bad:
        raise ValueError(
            f"{green_path} {sheet}: invalid header {', '.join(bad)}; expected columns A..K with "
            "B=port, C=module, I=*Component, J=*port")


def _read_green_rows(ws, location: str, errors: list[str]) -> list[GreenRow]:
    rows: list[GreenRow] = []
    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        app_port, module, data_element, peer, peer_port, remark = (
            norm(row[index]) if index < len(row) else "" for index in (1, 2, 7, 8, 9, 10))
        if not any([app_port, module, peer, peer_port]):
            continue
        # I/J 同时为空时静默跳过；仅一列为空且备注未说明预留用途时报错。
        if not peer and not peer_port:
            continue
        if not peer or not peer_port:
            if not _is_reserved_remark(remark):
                errors.append(f"{location} row {excel_row}: column I or J is empty")
            continue
        if not app_port:
            errors.append(f"{location} row {excel_row}: column B is empty")
            continue
        rows.append(GreenRow(app_port, module, peer, peer_port, data_element, remark, excel_row))
    return rows


def parse_green_workbook(green_path: Path, app_name: str) -> GreenWorkbookInfo:
    """读取单个绿表的 INPUT/OUTPUT 页；逐行记录全部错误后统一失败。

    按固定列号取值而不依赖表头文案，避免换行、空格、全角字符差异导致匹配失败。
    """

    workbook = load_workbook(green_path, read_only=True, data_only=True)
    try:
        if INPUT_SHEET not in workbook.sheetnames or OUTPUT_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Green table must contain {INPUT_SHEET}/{OUTPUT_SHEET} sheets: {green_path}")
        errors: list[str] = []
        rows: dict[str, list[GreenRow]] = {}
        for sheet in (INPUT_SHEET, OUTPUT_SHEET):
            _check_green_header(workbook[sheet], green_path, sheet)
            rows[sheet] = _read_green_rows(workbook[sheet], f"{green_path} {sheet}", errors)
    finally:
        workbook.close()
    for message in errors:
        LOGGER.error("%s", message)
    if errors:
        raise ValueError(f"{green_path}: {len(errors)} invalid row(s)")
    return GreenWorkbookInfo(app_name, rows[INPUT_SHEET], rows[OUTPUT_SHEET])


def collect_green_workbooks(
    entries: Iterable[str],
    variant: str,
    base_dir: Path,
) -> list[tuple[str, Path]]:
    """展开绿表路径和 glob，以文件名 stem 作为 app 名并去重。"""

    result: list[tuple[str, Path]] = []
    seen_paths: set[Path] = set()
    for entry in entries:
        rendered = norm(entry).replace("${variant}", variant)
        if not rendered:
            continue
        pattern = filesystem_path(rendered)
        if not pattern.is_absolute():
            pattern = base_dir / pattern

        if glob.has_magic(str(pattern)):
            # 通配符未命中视为没有可选绿表，不当作配置错误。
            matches = [Path(match) for match in glob.glob(str(pattern), recursive=True)]
        else:
            matches = [pattern]

        for match in sorted(matches, key=lambda path: str(path).casefold()):
            path = match.resolve()
            if path.name.startswith("~$"):
                continue
            if path.suffix.casefold() != ".xlsx" or not path.is_file():
                raise FileNotFoundError(f"Green table does not exist: {path}")
            if path not in seen_paths:
                seen_paths.add(path)
                result.append((path.stem, path))

    if not result:
        raise ValueError("No green tables were selected")
    return result


def _merged_sheet_name(app_name: str, sheet: str) -> str:
    return f"{app_name}_{sheet}"[:31]


def _copy_green_sheet(target_sheet, source_sheet) -> None:
    """复制单个绿表页签的内容、单元格样式和工作表格式。"""

    for (row_index, column_index), source_cell in source_sheet._cells.items():
        target_cell = target_sheet.cell(row=row_index, column=column_index)
        target_cell._value = source_cell._value
        target_cell.data_type = source_cell.data_type
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)
        target_cell.quotePrefix = source_cell.quotePrefix
        target_cell.pivotButton = source_cell.pivotButton
        if source_cell.hyperlink:
            target_cell._hyperlink = copy(source_cell.hyperlink)
        if source_cell.comment:
            target_cell.comment = copy(source_cell.comment)

    for key, dimension in source_sheet.row_dimensions.items():
        target_sheet.row_dimensions[key] = copy(dimension)
        target_sheet.row_dimensions[key].worksheet = target_sheet
    for key, dimension in source_sheet.column_dimensions.items():
        target_sheet.column_dimensions[key] = copy(dimension)
        target_sheet.column_dimensions[key].worksheet = target_sheet

    for merged_range in source_sheet.merged_cells.ranges:
        target_sheet.merge_cells(str(merged_range))

    target_sheet.sheet_format = copy(source_sheet.sheet_format)
    target_sheet.sheet_properties = copy(source_sheet.sheet_properties)
    target_sheet.page_margins = copy(source_sheet.page_margins)
    target_sheet.page_setup = copy(source_sheet.page_setup)
    target_sheet.print_options = copy(source_sheet.print_options)
    target_sheet.protection = copy(source_sheet.protection)
    target_sheet.views = deepcopy(source_sheet.views)
    target_sheet.freeze_panes = source_sheet.freeze_panes
    target_sheet.auto_filter = deepcopy(source_sheet.auto_filter)
    target_sheet.data_validations = deepcopy(source_sheet.data_validations)
    target_sheet.conditional_formatting = deepcopy(source_sheet.conditional_formatting)
    target_sheet.oddHeader = copy(source_sheet.oddHeader)
    target_sheet.oddFooter = copy(source_sheet.oddFooter)
    target_sheet.evenHeader = copy(source_sheet.evenHeader)
    target_sheet.evenFooter = copy(source_sheet.evenFooter)
    target_sheet.firstHeader = copy(source_sheet.firstHeader)
    target_sheet.firstFooter = copy(source_sheet.firstFooter)


def merge_green_workbooks(
    green_workbooks: list[tuple[str, Path]],
    output_file: Path,
) -> Path:
    """把本次处理到的所有绿表合并成一个工作簿，便于统一检索。"""

    output_file.parent.mkdir(parents=True, exist_ok=True)

    merged_workbook = Workbook()
    placeholder_sheet = merged_workbook.active
    if placeholder_sheet is None:
        raise RuntimeError("Cannot access default sheet in merged workbook")
    placeholder_sheet.title = "MergedGreenTables"

    for index, (app_name, green_workbook_path) in enumerate(green_workbooks):
        source_workbook = load_workbook(green_workbook_path, read_only=False, data_only=False)
        if index == 0:
            input_sheet = placeholder_sheet
            input_sheet.title = _merged_sheet_name(app_name, INPUT_SHEET)
        else:
            input_sheet = merged_workbook.create_sheet(_merged_sheet_name(app_name, INPUT_SHEET))

        output_sheet = merged_workbook.create_sheet(_merged_sheet_name(app_name, OUTPUT_SHEET))
        _copy_green_sheet(input_sheet, source_workbook[INPUT_SHEET])
        _copy_green_sheet(output_sheet, source_workbook[OUTPUT_SHEET])
        source_workbook.close()

    merged_workbook.save(output_file)
    merged_workbook.close()
    return output_file
