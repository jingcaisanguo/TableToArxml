"""Check an installed wheel outside the checkout, including real conversion baselines."""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def workbook_values(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook}
    finally:
        workbook.close()


def verify(python: Path) -> None:
    environment = dict(os.environ)
    environment.pop("PYTHONPATH", None)
    with TemporaryDirectory(prefix="tta-wheel-check-") as directory:
        temporary = Path(directory)

        def run(*arguments: object, expected: int = 0) -> str:
            result = subprocess.run(
                list(map(str, arguments)), cwd=temporary, env=environment,
                capture_output=True, text=True, encoding="utf-8", errors="replace",
            )
            if result.returncode != expected:
                raise RuntimeError(
                    f"Command {arguments!r} returned {result.returncode}, expected {expected}\n"
                    f"{result.stdout}\n{result.stderr}"
                )
            return result.stdout

        installed = json.loads(run(
            python, "-c",
            "import json, sysconfig, table_to_arxml; "
            "from importlib.metadata import version; "
            "from table_to_arxml.common.dependencies import check_dependencies; check_dependencies(); "
            "assert table_to_arxml.__version__ == version('table-to-arxml'); "
            "print(json.dumps({'file': table_to_arxml.__file__, 'scripts': sysconfig.get_path('scripts')}))",
        ))
        if ROOT / "src" in Path(installed["file"]).resolve().parents:
            raise RuntimeError(f"Expected an installed wheel, not checkout sources: {installed['file']}")
        scripts = Path(installed["scripts"])
        for tool in ("system", "swc"):
            command = scripts / f"table-to-{tool}-arxml{'.exe' if os.name == 'nt' else ''}"
            run(command, "--help")
            run(python, "-m", f"table_to_arxml.{tool}", "--help")
        run(scripts / f"table-to-system-arxml{'.exe' if os.name == 'nt' else ''}", expected=2)
        run(scripts / f"table-to-swc-arxml{'.exe' if os.name == 'nt' else ''}", expected=2)

        for label, tool, config, baseline in (
            ("bcm", "system", "BcmTableToSystemArxmlConfig.toml", "BcmTest"),
            ("mc", "system", "McTableToSystemArxmlConfig.toml", "McTest"),
            ("swc", "swc", "TableToSwcArxml.toml", ""),
        ):
            sample = ROOT / "test" / "samples" / tool
            output = temporary / label
            run(
                python, "-m", f"table_to_arxml.{tool}", "--config", sample / config,
                "--output-dir", output, "--log-dir", temporary / f"{label}-log",
            )
            expected_dir = sample / "expected" / baseline
            produced = {path.name: path for path in output.rglob("*") if path.is_file()}
            expected_files = {path.name: path for path in expected_dir.iterdir() if path.is_file()}
            if set(produced) != set(expected_files):
                raise RuntimeError(f"{label}: unexpected output files: {sorted(produced)}")
            for name, expected in expected_files.items():
                actual = produced[name]
                equal = workbook_values(actual) == workbook_values(expected) if actual.suffix == ".xlsx" else (
                    actual.read_bytes().replace(b"\r\n", b"\n") == expected.read_bytes().replace(b"\r\n", b"\n")
                )
                if not equal:
                    raise RuntimeError(f"{label}: output differs from baseline: {name}")
            print(f"Installed wheel: {label} baseline passed")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--python", type=Path, default=Path(sys.executable), help="Interpreter with the wheel installed")
    arguments = parser.parse_args()
    verify(arguments.python.absolute())


if __name__ == "__main__":
    main()
