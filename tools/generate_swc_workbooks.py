"""生成 SWC 空白模板和回归样例；两者共用列定义、填写说明和格式。"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from math import ceil
from pathlib import Path
from unicodedata import east_asian_width

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.views import Selection
from openpyxl.workbook.defined_name import DefinedName

from table_to_arxml.swc.schema import (
    ACCESS_KIND_CHOICES,
    ARRAY_SIZE_SEMANTICS,
    BOOLEAN_VALUES,
    CS_ARGUMENT_DIRECTIONS,
    SHEET_HEADERS,
)

ROOT = Path(__file__).resolve().parents[1]
BOOLEAN_CHOICES = BOOLEAN_VALUES
ACCESS_CHOICES = ACCESS_KIND_CHOICES
HEADER_COLORS = {"required": "1F4E78", "conditional": "9C6500", "optional": "516272"}


@dataclass(frozen=True)
class Field:
    name: str
    required: str
    default: str
    description: str
    example: str
    width: int = 28
    choices: tuple[str, ...] = ()
    number: str = ""

    @property
    def color(self) -> str:
        kind = "required" if self.required == "必填" else (
            "optional" if self.required == "可选" else "conditional"
        )
        return HEADER_COLORS[kind]


SWC = Field("SW-C", "必填", "—", "工作簿内组件标识；在 Components 中定义，不要求等于组件 SHORT-NAME。", "LIT", 14)
TARGET = Field("Target ARXML", "必填", "—", "要插入对象的已有文件，相对 projectRoot；不是 outputDir 下的导出文件名。所在包也须已存在。", "arxml/Types.arxml", 46)
TYPE_PATH = Field("Type Path", "必填", "—", "新建实现数据类型的完整 AUTOSAR 路径；同路径的成员/枚举行按顺序合并。", "/Pkg/Types/Mode", 46)
INTERFACE_PATH = Field("Interface Path", "必填", "—", "新建接口的完整 AUTOSAR 路径；同路径的多行合并。", "/Pkg/Interfaces/Status", 46)
SERVICE = Field("Is Service", "可选", "N", "接口级 IS-SERVICE；同一接口的所有行必须一致。", "Y / N", 16, BOOLEAN_CHOICES)
PORT_NAME = Field("Port Name", "必填", "—", "端口 SHORT-NAME；同一组件内 P/R 两页之间也不能重名。", "p_Status / r_Status", 28)
PORT_INTERFACE = Field("Port Interface Ref", "必填", "—", "已有或本次新建接口的完整路径。", "/Pkg/Interfaces/Status", 46)
ARTIFACT = Field("Create RTE Artifact", "可选", "N", "Y：SR 按 Access Kind 创建全部元素的访问点；CS P 创建 runnable/event，CS R 创建同步调用点。不生成 C 源码。", "Y / N", 22, BOOLEAN_CHOICES)
P_RUNNABLE = Field("Runnable Name", "SR 且 Artifact=Y 时必填", "CS 自动命名", "SR：已有 runnable 的简称。CS P：新建名称；留空用 RE_<端口>_<operation>，多 operation 时给手填名追加后缀。", "Run / RE_Service", 28)
R_RUNNABLE = Field("Runnable Name", "Artifact=Y 时必填", "—", "组件内部行为中已有 runnable 的简称，不是完整 AUTOSAR 路径。", "Run", 28)
P_ACCESS = Field("Access Kind", "可选", "Send", "仅 SR 且 Artifact=Y 生效；Send 为显式发送，Write 为隐式写访问。", "Send / Write", 26, ("Send", "Write"))
R_ACCESS = Field("Access Kind", "可选", "ReceiveByArgument", "仅 SR 且 Artifact=Y 生效；ReceiveBy* 为显式接收，Read 为隐式读访问。", "ReceiveByValue / Read", 26, ("ReceiveByArgument", "ReceiveByValue", "Read"))

LAYOUT: dict[str, tuple[Field, ...]] = {
    "Components": (
        SWC,
        Field("Component Path", "必填", "—", "清单中已有的软件组件类型路径；本页只建立映射，不新建组件。", "/BcmSystem/SwComponentTypes/LIT", 54),
    ),
    "Enumerations": (
        TYPE_PATH,
        Field("Value", "必填", "—", "整数枚举值；同一类型内不能重复。", "0", 14, number="integer"),
        Field("Label", "必填", "—", "枚举值对应文本；同一类型内不能重复。", "Mode_Off", 28),
        Field("Implementation Data Type Ref", "必填", "—", "底层 IMPLEMENTATION-DATA-TYPE 的路径，不是 SW-BASE-TYPE。", "/AUTOSAR_Platform/ImplementationDataTypes/uint8", 46),
        Field("Compu Method Ref", "可选", "按 Type Path 推导", "生成的计算方法路径；默认同级 CompuMethods 包中的同名对象。", "/Pkg/CompuMethods/Mode", 40),
        Field("Data Constr Ref", "可选", "按 Type Path 推导", "生成的数据约束路径；默认同级 DataConstrs 包中的 dc_<类型名>。", "/Pkg/DataConstrs/dc_Mode", 40),
        TARGET,
    ),
    "Arrays": (
        TYPE_PATH,
        Field("Element Name", "必填", "—", "数组实现数据类型中的元素 SHORT-NAME。", "Element", 28),
        Field("Element Type Ref", "必填", "—", "数组元素的实现数据类型路径；引用其他元类（如 SW-BASE-TYPE）报错。", "/Pkg/Types/Mode", 46),
        Field("Array Size", "必填", "—", "正整数；VARIABLE-SIZE 时表示声明的最大长度。", "8", 16, number="positive_integer"),
        Field("Array Size Semantics", "可选", "FIXED-SIZE", "数组大小语义。", "FIXED-SIZE / VARIABLE-SIZE", 26, ARRAY_SIZE_SEMANTICS),
        TARGET,
    ),
    "Structures": (
        TYPE_PATH,
        Field("Member Name", "必填", "—", "一个成员一行；同一结构体成员名唯一，按行顺序生成。", "State", 28),
        Field("Member Type Ref", "必填", "—", "成员的实现数据类型路径，可引用本次新建的类型。", "/Pkg/Types/Mode", 46),
        TARGET,
    ),
    "SR-Interfaces": (
        INTERFACE_PATH,
        Field("Data Element Name", "必填", "—", "一个数据元素一行；同一接口内名称唯一。", "State", 28),
        Field("Data Type Ref", "必填", "—", "数据元素的数据类型路径（实现或应用数据类型），DEST 取被引用对象的元类。", "/Pkg/Types/Mode", 46),
        Field("Is Queued", "可选", "N", "Y 时生成 SW-IMPL-POLICY=QUEUED。", "Y / N", 16, BOOLEAN_CHOICES),
        SERVICE, TARGET,
    ),
    "CS-Interfaces": (
        INTERFACE_PATH,
        Field("Operation Name", "必填", "—", "一个参数一行；同一接口、同一 operation 的参数按行合并。无参数 operation 只填一行。", "Reset / SetMode", 28),
        Field("Argument Name", "有参数时必填", "全空表示无参数", "参数名、方向和类型必须全填或全空；同一 operation 不得混用无参数声明和参数行。", "Mode", 28),
        Field("Direction", "有参数时必填", "—", "参数方向；与 Argument Name、Data Type Ref 一起填写。", "IN / OUT / INOUT", 16, CS_ARGUMENT_DIRECTIONS),
        Field("Data Type Ref", "有参数时必填", "—", "参数的数据类型路径（实现或应用数据类型）；无参数 operation 留空。", "/Pkg/Types/Mode", 46),
        SERVICE, TARGET,
    ),
    "P-Ports": (
        SWC, PORT_NAME, PORT_INTERFACE, ARTIFACT, P_RUNNABLE, P_ACCESS,
        Field("Transmission Acknowledge Timeout", "可选", "不生成", "SR 发送确认超时，单位秒；非负数，填写 0 也会触发 ComSpec 生成。", "0.01", 28, number="nonnegative"),
        Field("Uses End To End Protection", "可选", "N", "SR 发送端 E2E 标志；仅填 N 不会触发 ComSpec 生成。", "Y / N", 26, BOOLEAN_CHOICES),
        Field("Runnable Symbol", "可选", "与新建 runnable 名相同", "仅 CS P 且 Artifact=Y 生效；C 函数符号，多 operation 时追加后缀。", "Service_Run", 28),
        Field("Minimum Start Interval", "可选", "0.0 秒", "仅 CS P 且 Artifact=Y 生效；非负最小启动间隔，单位秒。", "0.01", 26, number="nonnegative"),
        Field("Can Be Invoked Concurrently", "可选", "N", "仅 CS P 且 Artifact=Y 生效；允许并发调用。", "Y / N", 28, BOOLEAN_CHOICES),
    ),
    "R-Ports": (
        SWC, PORT_NAME, PORT_INTERFACE, ARTIFACT, R_RUNNABLE, R_ACCESS,
        Field("Enable Update", "可选", "N", "Y 时为非队列数据元素创建 NONQUEUED 接收 ComSpec（ENABLE-UPDATE）；全部元素为队列元素时报错。", "Y / N", 20, BOOLEAN_CHOICES),
        Field("Queue Length", "可选", "0 / 不生成", "非负整数；正数为队列元素（SW-IMPL-POLICY=QUEUED）生成 QUEUED 接收 ComSpec，没有队列元素时报错；0 或留空不生成。", "4", 20, number="nonnegative_integer"),
    ),
    "Data Accesses": (
        SWC,
        Field("Runnable Name", "必填", "—", "已有或本次 CS P 新建 runnable 的简称；缺少访问点容器时自动创建。", "Run", 28),
        Field("Port Name", "必填", "—", "已有或本次新建端口的 SHORT-NAME。", "r_Status", 28),
        Field("Data Element Name", "多元素接口时必填", "单元素自动选择", "只处理指定的数据元素。", "State", 28),
        Field("Access Kind", "可选", "P: Send / R: ReceiveByArgument", "须匹配端口方向；本页可自定义名称或仅处理部分元素。", " / ".join(ACCESS_CHOICES), 26, ACCESS_CHOICES),
        Field("Access Name", "可选", "按端口名自动生成", "访问点 SHORT-NAME；同一 runnable 内不能冲突。", "dra_Status", 28),
    ),
}

SAMPLE_ROWS = {
    "Components": [("LIT", "/BcmSystem/SwComponentTypes/LIT"), ("Horn", "/BcmSystem/SwComponentTypes/Horn")],
    "P-Ports": [
        ("LIT", "p_LIT_TestPowerMode", "/BcmTest/PowerModeInterface", "N"),
        ("Horn", "p_Horn_TestPositionActive", "/BcmTest/PositionActiveInterface", "N"),
    ],
    "R-Ports": [
        ("LIT", "r_LIT_TestPositionActive", "/BcmTest/PositionActiveInterface", "N"),
        ("Horn", "r_Horn_TestPowerMode", "/BcmTest/PowerModeInterface", "N", None, None, "Y"),
    ],
}


def display_width(text: str) -> int:
    return sum(2 if east_asian_width(char) in "WF" else 1 for char in text)


def row_height(values: list[str], widths: list[int]) -> float:
    lines = max(
        sum(max(1, ceil(display_width(line) / (width - 3))) for line in value.split("\n"))
        for value, width in zip(values, widths)
    )
    return max(32, 16 * lines + 12)


def link(cell, sheet: str, target: str) -> None:
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{target}")


def add_validation(sheet, column: str, field: Field) -> None:
    formula = ""
    kind = "list"
    if field.name == "SW-C" and sheet.title != "Components":
        formula = "=SwcComponentKeys"
    elif field.choices:
        formula = '"' + ",".join(field.choices) + '"'
    elif field.number:
        kind = "custom"
        ref = f"{column}2"
        number = f"VALUE({ref})"
        checks = []
        if "integer" in field.number:
            checks.append(f"MOD({number},1)=0")
        if field.number.startswith("positive"):
            checks.append(f"{number}>0")
        elif field.number.startswith("nonnegative"):
            checks.append(f"{number}>=0")
        formula = f'OR({ref}="",IFERROR(AND({",".join(checks)}),FALSE))'
    if not formula:
        return
    validation = DataValidation(
        type=kind, formula1=formula, allow_blank=True, showDropDown=False,
        showErrorMessage=True, showInputMessage=True, errorStyle="stop",
        promptTitle="填写提示", prompt=f"{field.required}；默认：{field.default}。{field.description}",
        errorTitle="输入不符合本列规则",
        error="请按表头批注填写。粘贴可能绕过 Excel 验证，运行工具时仍会完整校验。",
    )
    sheet.add_data_validation(validation)
    validation.add(f"{column}2:{column}1048576")


def make_readme(workbook: Workbook, sample: bool) -> dict[str, int]:
    sheet = workbook.active
    sheet.title = "README"
    widths = [24, 36, 28, 28, 64, 46]
    for column, width in zip("ABCDEF", widths):
        sheet.column_dimensions[column].width = width
    sheet.sheet_properties.tabColor = "1F4E78"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.selection = [Selection(activeCell="A1", sqref="A1")]
    sheet.sheet_view.topLeftCell = "A1"
    sheet.freeze_panes = None
    sheet.auto_filter.ref = None

    def paragraph(title: str, text: str, color: str = "EAF0F5") -> None:
        row = sheet.max_row + 1 if sheet["A1"].value else 1
        sheet.cell(row, 1, title)
        sheet.cell(row, 2, text)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        for cell in sheet[row]:
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(name="Calibri", size=11, bold=cell.column == 1, color="243238")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row].height = max(36, row_height([title, text], [widths[0], sum(widths[1:])]))

    paragraph("SWC ARXML 工作簿", "测试样例：包含 LIT/Horn 数据，不要直接用于其他项目。" if sample else "空白模板：填写后另存，并在配置中指定 workbook 路径。")
    paragraph("用途与边界", "向已有 ARXML 增量添加类型、接口、端口和 RTE 访问对象；不新建组件、不修改已有对象、不生成 C 源码。")
    paragraph("流程 1：只加端口", "Components → P-Ports / R-Ports。前三列必填，Create RTE Artifact 留空或 N 即仅创建端口及选配 ComSpec。")
    paragraph("流程 2：新接口", "类型（需要时）→ SR-Interfaces / CS-Interfaces → P/R-Ports。端口可以引用本次新建接口，接口可以引用本次新建类型。")
    paragraph("流程 3：访问点", "端口页 Artifact=Y：为全部元素自动创建；Data Accesses：只选部分元素或自定义名称。同一访问需求不要在两处重复填写。")
    paragraph("页签与表头", "未使用的工作表可以省略；有端口或访问点时必须提供 Components 映射。只接受本模板规定的页名和列名（不区分大小写），其他页名/列名报错；README 仅供阅读。第 1 行是表头，禁止增加标题行、空白分隔列或给列名加星号。")
    paragraph("颜色与位置", "深蓝表头=必填，棕色=条件必填，灰蓝=可选；具体条件见批注。P/R 前六列位置一致，后面为 ComSpec 和 CS 服务端配置。点击各页 A1 返回本页说明。")
    paragraph("路径与名称", "Path 是对象完整 AUTOSAR 路径（组件路径指已有对象）；Ref 是引用路径；Name 是 SHORT-NAME。SW-C 是工作簿标识。AUTOSAR 路径以 / 开头。")
    paragraph("名称校验", "新建 Name 和路径各段必须以 ASCII 字母开头，只含字母、数字、下划线，最长 128 字符。同一命名空间内名称不区分大小写唯一。枚举 Label 可用中文，但非 C 标识符会 WARNING（RTE 可能需要 SYMBOL）；Runnable Symbol 使用 C 标识符。自动长名称截断并追加哈希。")
    paragraph("目标文件", "Target ARXML 相对 projectRoot，目标文件和所在包必须已存在；它不是导出路径。最终输出由 outputDir / inPlace / dryRun 控制。")
    paragraph("输出与重定义", "outputDir 不能覆盖任何输入 ARXML（包括链接别名），覆盖源文件必须明确使用 inPlace。其他文件已有同路径对象时禁止再次创建；候选文件解析失败也会阻止写出。")
    paragraph("空值与布尔值", "布尔留空为 N；接受 Y/N、TRUE/FALSE、YES/NO、1/0（不区分大小写）。其他默认值逐列列出。合法但不生效的配置会 WARNING；非法值仍 ERROR。")
    paragraph("错误与写出", "任一 ERROR 阻止全部 ARXML 写出，退出码 1；WARNING 不阻止写出。空模板不产生对象，运行会提示没有有效记录。")
    paragraph("数据表操作", "每页使用 Excel 表格：从表格最后一个单元格按 Tab 追加行，筛选会自动扩展。若在表外粘贴，使用“调整表格大小”纳入新行；解析器仍读取整页。")
    paragraph("排序与合并", "批量数据优先逐行重复公共字段；合并单元格虽可读取，但不利于 Excel 表格排序/筛选，不建议在数据表中合并。")
    paragraph("列移动与下拉", "工具按列名读取；移动整列后应检查 Excel 数据验证范围。下拉、数值验证和批注用于辅助，复制粘贴仍须经工具校验。")
    paragraph("SR P / SR R", "Artifact=Y 时必须指定已有 Runnable Name，缺少的访问点容器自动创建。P：Send（默认）/ Write；R：ReceiveByArgument（默认）/ ReceiveByValue / Read；队列元素不能用 ReceiveByValue / Read，Enable Update=Y 时不能用 Read。")
    paragraph("CS P / CS R", "CS P：每个 operation 新建 runnable/event；名称留空自动命名，多 operation 给手填名追加后缀。CS R：已有调用方 Runnable Name 必填，每个 operation 创建同步调用点（超时 0.0）。")
    paragraph("ComSpec 触发", "P：超时非空（含 0）或 E2E=Y，按元素 SW-IMPL-POLICY 生成 QUEUED / NONQUEUED 发送 ComSpec。R：Enable Update=Y 作用于非队列元素，Queue Length>0 作用于队列元素，不适用于任何元素时报错。仅填 N 或 Queue Length=0 不触发。")
    paragraph("已知限制", "组件、内部行为和 PORTS 容器须预先存在，访问点容器自动创建；不支持 PR-Port；CS 不生成应用错误；SR 接口以外不自动生成数据访问点。不接受公式和含换行/制表符的单元格。无参数 operation 的参数三列全空，不能混用同名有参数声明。")
    paragraph("精确类型", "枚举底层列引用 IMPLEMENTATION-DATA-TYPE，不是 SW-BASE-TYPE。所有时间间隔/超时单位为秒。")
    paragraph("字段导航", "点击下方工作表标题跳转输入页；字段说明与各页表头批注来自同一份定义。")
    sections = {}
    for name, fields in LAYOUT.items():
        paragraph(name, "字段、必填条件、默认值、说明和示例如下。", "D9EAF7")
        sections[name] = sheet.max_row
        link(sheet.cell(sheet.max_row, 1), name, "A1")
        sheet.append(["工作表", "字段", "必填条件", "默认值", "说明", "示例 / 取值"])
        for cell in sheet[sheet.max_row]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[sheet.max_row].height = 30
        for field in fields:
            values = [name, field.name, field.required, field.default, field.description, field.example]
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                cell.font = Font(name="Calibri", size=11, color="243238")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.fill = PatternFill("solid", fgColor="F0F5F6" if sheet.max_row % 2 else "FFFFFF")
            sheet.row_dimensions[sheet.max_row].height = row_height(values, widths)
    return sections


def create_workbook(sample: bool = False) -> Workbook:
    if set(LAYOUT) != set(SHEET_HEADERS):
        raise ValueError("Template worksheets do not match the reader schema")
    for name, fields in LAYOUT.items():
        required, optional = SHEET_HEADERS[name]
        if {field.name for field in fields} != set(required + optional) or len(fields) != len(required + optional):
            raise ValueError(f"Template columns do not match the reader schema: {name}")
        if any((field.name in required) != (field.required == "必填") for field in fields):
            raise ValueError(f"Template required columns do not match the reader schema: {name}")
    workbook = Workbook()
    sections = make_readme(workbook, sample)
    tab_colors = {
        "Components": "1F4E78", "Enumerations": "247B7B", "Arrays": "247B7B",
        "Structures": "247B7B", "SR-Interfaces": "7355A0", "CS-Interfaces": "7355A0",
        "P-Ports": "B36B18", "R-Ports": "B36B18", "Data Accesses": "387D42",
    }
    for name, fields in LAYOUT.items():
        sheet = workbook.create_sheet(name)
        sheet.append([field.name for field in fields])
        for values in SAMPLE_ROWS.get(name, ()) if sample else ():
            sheet.append(values)
        if sheet.max_row == 1:
            sheet.append([None] * len(fields))
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.topLeftCell = "A1"
        sheet.freeze_panes = "C2" if name in ("P-Ports", "R-Ports", "Data Accesses") else "A2"
        for selection in sheet.sheet_view.selection:
            selection.activeCell = "A2"
            selection.sqref = "A2"
        sheet.sheet_properties.tabColor = tab_colors[name]
        sheet.row_dimensions[1].height = 64
        for index, field in enumerate(fields, start=1):
            cell = sheet.cell(1, index)
            sheet.column_dimensions[cell.column_letter].width = field.width
            cell.fill = PatternFill("solid", fgColor=field.color)
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            cell.border = Border(left=Side(style="medium" if index in (4, 7, 9) and name == "P-Ports" else "thin", color="FFFFFF"))
            cell.comment = Comment(f"{field.required}\n默认：{field.default}\n{field.description}\n示例：{field.example}", "TableToArxml")
            cell.comment.width = 360
            cell.comment.height = 140
            add_validation(sheet, cell.column_letter, field)
            for row in range(2, sheet.max_row + 1):
                body = sheet.cell(row, index)
                body.font = Font(name="Calibri", size=11, color="243238")
                body.alignment = Alignment(vertical="top", wrap_text=True)
                body.border = Border(bottom=Side(style="hair", color="D9E2F3"))
                body.number_format = "0" if "integer" in field.number else (
                    "0.########" if field.number else "General" if field.choices else "@"
                )
        for row in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row].height = row_height(
                [str(sheet.cell(row, col).value or "") for col in range(1, len(fields) + 1)],
                [field.width for field in fields],
            )
        link(sheet["A1"], "README", f"A{sections[name]}")
        table = Table(displayName="Swc" + name.replace("-", "").replace(" ", ""), ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    workbook.defined_names.add(DefinedName("SwcComponentKeys", attr_text="SwcComponents[SW-C]"))
    for sheet in workbook:
        sheet.print_title_rows = "1:2" if sheet.title == "README" else "1:1"
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.active = 0
    return workbook


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, help="验证用输出目录；不指定时更新仓库内两份工作簿")
    args = parser.parse_args()
    outputs = (
        ("SwcTestPorts.xlsx", ROOT / "test" / "samples" / "swc", True),
        ("SwcTemplate.xlsx", ROOT / "docs" / "user" / "swc", False),
    )
    for filename, directory, sample in outputs:
        target = (args.output_dir if args.output_dir is not None else directory) / filename
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook = create_workbook(sample)
        try:
            workbook.save(target)
        finally:
            workbook.close()
        print(target)


if __name__ == "__main__":
    main()
