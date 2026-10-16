"""System 工具测试用的最小 ARXML 模型和白表构造函数。"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

MODEL = """<?xml version="1.0" encoding="UTF-8"?>
<AUTOSAR xmlns="http://autosar.org/schema/r4.0" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:schemaLocation="http://autosar.org/schema/r4.0 AUTOSAR_4-3-0.xsd">
  <AR-PACKAGES>
    <AR-PACKAGE>
      <SHORT-NAME>Pkg</SHORT-NAME>
      <ELEMENTS>
        <SW-BASE-TYPE>
          <SHORT-NAME>u8</SHORT-NAME>
          <CATEGORY>FIXED_LENGTH</CATEGORY>
          <BASE-TYPE-SIZE>8</BASE-TYPE-SIZE>
          <BASE-TYPE-ENCODING>NONE</BASE-TYPE-ENCODING>
        </SW-BASE-TYPE>
        <SW-BASE-TYPE>
          <SHORT-NAME>u16</SHORT-NAME>
          <CATEGORY>FIXED_LENGTH</CATEGORY>
          <BASE-TYPE-SIZE>16</BASE-TYPE-SIZE>
          <BASE-TYPE-ENCODING>NONE</BASE-TYPE-ENCODING>
        </SW-BASE-TYPE>
{implementation_types}
{interfaces}
        <APPLICATION-SW-COMPONENT-TYPE>
          <SHORT-NAME>A</SHORT-NAME>
          <PORTS>
            <P-PORT-PROTOTYPE>
              <SHORT-NAME>pOut</SHORT-NAME>
              <PROVIDED-INTERFACE-TREF DEST="SENDER-RECEIVER-INTERFACE">/Pkg/IfA</PROVIDED-INTERFACE-TREF>
            </P-PORT-PROTOTYPE>
            <P-PORT-PROTOTYPE>
              <SHORT-NAME>pWide</SHORT-NAME>
              <PROVIDED-INTERFACE-TREF DEST="SENDER-RECEIVER-INTERFACE">/Pkg/IfC</PROVIDED-INTERFACE-TREF>
            </P-PORT-PROTOTYPE>
          </PORTS>
        </APPLICATION-SW-COMPONENT-TYPE>
        <APPLICATION-SW-COMPONENT-TYPE>
          <SHORT-NAME>B</SHORT-NAME>
          <PORTS>
            <R-PORT-PROTOTYPE>
              <SHORT-NAME>rIn</SHORT-NAME>
              <REQUIRED-INTERFACE-TREF DEST="SENDER-RECEIVER-INTERFACE">/Pkg/IfB</REQUIRED-INTERFACE-TREF>
            </R-PORT-PROTOTYPE>
          </PORTS>
        </APPLICATION-SW-COMPONENT-TYPE>
{compositions}
        <SYSTEM>
          <SHORT-NAME>Sys</SHORT-NAME>
          <CATEGORY>SYSTEM_DESCRIPTION</CATEGORY>
          <ROOT-SOFTWARE-COMPOSITIONS>
            <ROOT-SW-COMPOSITION-PROTOTYPE>
              <SHORT-NAME>RootProto</SHORT-NAME>
              <SOFTWARE-COMPOSITION-TREF DEST="COMPOSITION-SW-COMPONENT-TYPE">/Pkg/Root</SOFTWARE-COMPOSITION-TREF>
            </ROOT-SW-COMPOSITION-PROTOTYPE>
          </ROOT-SOFTWARE-COMPOSITIONS>
        </SYSTEM>
      </ELEMENTS>
    </AR-PACKAGE>
    <AR-PACKAGE>
      <SHORT-NAME>Other</SHORT-NAME>
      <ELEMENTS>
        <APPLICATION-SW-COMPONENT-TYPE>
          <SHORT-NAME>A</SHORT-NAME>
        </APPLICATION-SW-COMPONENT-TYPE>
      </ELEMENTS>
    </AR-PACKAGE>
  </AR-PACKAGES>
</AUTOSAR>
"""


def _implementation_type(name: str, base: str) -> str:
    return f"""        <IMPLEMENTATION-DATA-TYPE>
          <SHORT-NAME>{name}</SHORT-NAME>
          <CATEGORY>VALUE</CATEGORY>
          <SW-DATA-DEF-PROPS>
            <SW-DATA-DEF-PROPS-VARIANTS>
              <SW-DATA-DEF-PROPS-CONDITIONAL>
                <BASE-TYPE-REF DEST="SW-BASE-TYPE">/Pkg/{base}</BASE-TYPE-REF>
              </SW-DATA-DEF-PROPS-CONDITIONAL>
            </SW-DATA-DEF-PROPS-VARIANTS>
          </SW-DATA-DEF-PROPS>
        </IMPLEMENTATION-DATA-TYPE>"""


def _interface(name: str, element: str, data_type: str) -> str:
    return f"""        <SENDER-RECEIVER-INTERFACE>
          <SHORT-NAME>{name}</SHORT-NAME>
          <IS-SERVICE>false</IS-SERVICE>
          <DATA-ELEMENTS>
            <VARIABLE-DATA-PROTOTYPE>
              <SHORT-NAME>{element}</SHORT-NAME>
              <TYPE-TREF DEST="IMPLEMENTATION-DATA-TYPE">/Pkg/{data_type}</TYPE-TREF>
            </VARIABLE-DATA-PROTOTYPE>
          </DATA-ELEMENTS>
        </SENDER-RECEIVER-INTERFACE>"""


def _composition(name: str, prototypes: list[tuple[str, str, str]]) -> str:
    children = "\n".join(
        f"""            <SW-COMPONENT-PROTOTYPE>
              <SHORT-NAME>{prototype}</SHORT-NAME>
              <TYPE-TREF DEST="{dest}">{type_path}</TYPE-TREF>
            </SW-COMPONENT-PROTOTYPE>"""
        for prototype, dest, type_path in prototypes
    )
    return f"""        <COMPOSITION-SW-COMPONENT-TYPE>
          <SHORT-NAME>{name}</SHORT-NAME>
          <COMPONENTS>
{children}
          </COMPONENTS>
        </COMPOSITION-SW-COMPONENT-TYPE>"""


def model_text() -> str:
    application = "APPLICATION-SW-COMPONENT-TYPE"
    composition = "COMPOSITION-SW-COMPONENT-TYPE"
    return MODEL.format(
        implementation_types="\n".join(
            [_implementation_type("uint8", "u8"), _implementation_type("uint16", "u16")]
        ),
        interfaces="\n".join([
            _interface("IfA", "a", "uint8"),
            _interface("IfB", "b", "uint8"),
            _interface("IfC", "c", "uint16"),
        ]),
        compositions="\n".join([
            _composition("Top", [("A", application, "/Pkg/A"), ("B", application, "/Pkg/B")]),
            _composition("Root", [("TopProto", composition, "/Pkg/Top")]),
            _composition("Loop", [("Self", composition, "/Pkg/Loop")]),
        ]),
    )


def write_model(directory: Path) -> Path:
    """写出模型和只含该模型的清单，返回清单路径。"""

    (directory / "model.arxml").write_text(model_text(), encoding="utf-8")
    list_file = directory / "model.lst"
    list_file.write_text("model.arxml\n", encoding="utf-8")
    return list_file


def write_white_table(directory: Path, rows: dict[str, list[list[str]]]) -> Path:
    """按白表布局写出一个白表；``rows`` 为 ``页名 -> 数据行``。"""

    from table_to_arxml.system.white_table import create_white_workbook

    workbook = create_white_workbook({
        "Composition": "/Pkg/Top", "System": "/Pkg/Sys",
        "RootSwCompositionPrototype": "/Pkg/Sys/RootProto", "EcuInstance": "/Pkg/Ecu",
    })
    for sheet_name, sheet_rows in rows.items():
        for row in sheet_rows:
            workbook[sheet_name].append(row)
    path = directory / "white.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def write_config(directory: Path, **options: object) -> Path:
    config = {
        "defaultVariant": "T", "projectRoot": ".", "outputPrefix": "T", "whiteTable": "white.xlsx",
        "arxmlFileList": "model.lst", "outputDir": "out", "logDir": "log", "generateSystem": False,
        **options,
    }
    path = directory / "config.json"
    path.write_text(json.dumps(config), encoding="utf-8")
    return path


def sheet_rows(path: Path, sheet: str) -> list[tuple]:
    workbook = load_workbook(path, read_only=True)
    try:
        return list(workbook[sheet].iter_rows(values_only=True))
    finally:
        workbook.close()
