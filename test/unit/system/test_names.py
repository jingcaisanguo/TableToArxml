"""连接器命名与合并绿表页签名。"""

from __future__ import annotations

from openpyxl import Workbook, load_workbook
import pytest

from helpers import ROOT
from table_to_arxml.common.names import MAX_SHORT_NAME
from table_to_arxml.system.arxml_generator import connector_name, disambiguated_connector_name
from table_to_arxml.system.green_table import (
    MAX_SHEET_NAME,
    R_PORT_SHEET,
    P_PORT_SHEET,
    _merged_sheet_name,
    merge_green_workbooks,
)


def test_connector_name_separates_parts():
    assert connector_name("dlc_", "O", "AB", "C") == "dlc_O_AB_C"
    assert connector_name("dlc_", "O", "AB", "C") != connector_name("dlc_", "O", "A", "BC")


def test_ambiguous_connector_names_are_disambiguated_by_structured_hash():
    base = connector_name("ac_", "a_b", "c")
    assert base == connector_name("ac_", "a", "b_c")
    first = disambiguated_connector_name(base, "a_b", "c")
    second = disambiguated_connector_name(base, "a", "b_c")
    assert first != second
    assert first.startswith(base + "_") and second.startswith(base + "_")
    long_name = connector_name("ac_", "P" * 80, "port", "R" * 80, "port")
    assert len(disambiguated_connector_name(long_name, "P" * 80, "port")) == MAX_SHORT_NAME


def test_long_connector_names_are_truncated_with_stable_hash():
    first = connector_name("ac_", "P" * 80, "port", "R" * 80, "port")
    second = connector_name("ac_", "P" * 80, "port", "R" * 80, "other")
    assert len(first) == len(second) == MAX_SHORT_NAME
    assert first != second
    assert first == connector_name("ac_", "P" * 80, "port", "R" * 80, "port")


def test_merged_sheet_names_are_unique_within_excel_limit():
    used: set[str] = set()
    names = [
        _merged_sheet_name("VeryLongApplicationNameForTest" + suffix, R_PORT_SHEET, used)
        for suffix in ("A", "B", "C")
    ]
    assert len(set(name.casefold() for name in names)) == 3
    assert all(len(name) <= MAX_SHEET_NAME for name in names)
    assert names[0] == "VeryLongApplicationNameForTestA"[:MAX_SHEET_NAME]


def test_merge_green_workbooks_with_colliding_names(tmp_path):
    source = tmp_path / "green.xlsx"
    workbook = Workbook()
    for title in (R_PORT_SHEET, P_PORT_SHEET):
        sheet = workbook.active if title == R_PORT_SHEET else workbook.create_sheet()
        sheet.title = title
        sheet.append(["No", "port", "module"])
        sheet.merge_cells("A2:B2")
    workbook.save(source)
    workbook.close()

    merged = merge_green_workbooks(
        [("SameLongPrefixForApplication1", source), ("SameLongPrefixForApplication2", source)],
        tmp_path / "merged.xlsx",
    )
    result = load_workbook(merged)
    try:
        assert len(result.sheetnames) == len({name.casefold() for name in result.sheetnames}) == 4
        used_names: set[str] = set()
        expected_names = [
            _merged_sheet_name(app_name, sheet_name, used_names)
            for app_name in ("SameLongPrefixForApplication1", "SameLongPrefixForApplication2")
            for sheet_name in (R_PORT_SHEET, P_PORT_SHEET)
        ]
        assert result.sheetnames == expected_names
        assert all(str(ranges) == "A2:B2" for sheet in result for ranges in sheet.merged_cells.ranges)
    finally:
        result.close()


def test_sample_green_tables_merge(tmp_path):
    tables = sorted((ROOT / "test/samples/system/GreenTable/Bcm").glob("*.xlsx"))
    merged = merge_green_workbooks([(path.stem, path) for path in tables], tmp_path / "merged.xlsx")
    result = load_workbook(merged, read_only=True)
    try:
        assert len(result.sheetnames) == 2 * len(tables)
        assert result.sheetnames == [
            name for app_name, _ in [(path.stem, path) for path in tables]
            for name in (f"{app_name}_{R_PORT_SHEET}", f"{app_name}_{P_PORT_SHEET}")
        ]
    finally:
        result.close()


def test_parse_green_workbook_accepts_input_output_legacy_names(tmp_path):
    from table_to_arxml.system.green_table import parse_green_workbook

    source = ROOT / "test/samples/system/GreenTable/Bcm/LIT.xlsx"
    canonical = parse_green_workbook(source, "LIT")
    legacy_path = tmp_path / "legacy.xlsx"
    workbook = load_workbook(source)
    workbook[R_PORT_SHEET].title = "INPUT"
    workbook[P_PORT_SHEET].title = "OUTPUT"
    workbook.save(legacy_path)
    workbook.close()

    legacy = parse_green_workbook(legacy_path, "LIT")
    assert legacy.r_port_rows == canonical.r_port_rows
    assert legacy.p_port_rows == canonical.p_port_rows
    merged_path = merge_green_workbooks([("LIT", legacy_path)], tmp_path / "merged.xlsx")
    merged = load_workbook(merged_path, read_only=True)
    try:
        assert merged.sheetnames == ["LIT_R-Port", "LIT_P-Port"]
    finally:
        merged.close()


def test_parse_green_workbook_rejects_mixed_sheet_pairs(tmp_path):
    import pytest
    from table_to_arxml.system.green_table import parse_green_workbook

    source = ROOT / "test/samples/system/GreenTable/Bcm/LIT.xlsx"
    workbook = load_workbook(source)
    workbook["P-Port"].title = "INPUT"
    mixed_path = tmp_path / "mixed.xlsx"
    workbook.save(mixed_path)
    workbook.close()

    with pytest.raises(ValueError, match="exactly one complete"):
        parse_green_workbook(mixed_path, "LIT")


def test_parse_green_workbook_rejects_receive_send_aliases(tmp_path):
    from table_to_arxml.system.green_table import parse_green_workbook

    source = ROOT / "test/samples/system/GreenTable/Bcm/LIT.xlsx"
    workbook = load_workbook(source)
    workbook["R-Port"].title = "RECEIVE"
    workbook["P-Port"].title = "SEND"
    alias_path = tmp_path / "receive-send.xlsx"
    workbook.save(alias_path)
    workbook.close()

    with pytest.raises(ValueError, match="exactly one complete"):
        parse_green_workbook(alias_path, "LIT")
