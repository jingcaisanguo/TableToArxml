"""测试共用的路径、命令运行与基线比较函数。"""

from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
SYSTEM = "table_to_arxml.system"
SWC = "table_to_arxml.swc"
NS = {"ar": "http://autosar.org/schema/r4.0"}


def _environment() -> dict[str, str]:
    environment = dict(os.environ)
    environment["PYTHONPATH"] = os.pathsep.join(
        [str(SRC), *filter(None, [environment.get("PYTHONPATH")])]
    )
    return environment


def run_tool(module: str, *args: object, cwd: Path = ROOT) -> subprocess.CompletedProcess:
    """以 ``python -m`` 运行源码中的工具，等同安装后的命令。"""

    return subprocess.run(
        [sys.executable, "-m", module, *map(str, args)],
        cwd=cwd, env=_environment(), capture_output=True, text=True, encoding="utf-8", errors="replace",
    )


def run_python(code: str) -> subprocess.CompletedProcess:
    """在新进程中运行一段代码，用于检查模块导入副作用。"""

    return subprocess.run(
        [sys.executable, "-c", code],
        cwd=ROOT, env=_environment(), capture_output=True, text=True, encoding="utf-8", errors="replace",
    )


def cells(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}
    finally:
        workbook.close()


def lf_bytes(path: Path) -> bytes:
    # 基线和检出的样例可能是 CRLF 或 LF，比较前统一换行符。
    return path.read_bytes().replace(b"\r\n", b"\n")


def assert_matches_baseline(actual: Path, expected: Path) -> None:
    assert actual.is_file(), f"missing output {actual}"
    if expected.suffix == ".xlsx":
        assert cells(actual) == cells(expected), f"{actual.name} differs from {expected}"
    else:
        assert lf_bytes(actual) == lf_bytes(expected), f"{actual.name} differs from {expected}"
