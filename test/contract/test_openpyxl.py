"""openpyxl 契约测试：跨工作簿复制页签依赖的私有属性及复制结果。"""

from __future__ import annotations

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from table_to_arxml.common.openpyxl_compat import copy_worksheet


def test_private_attributes_used_for_copy_exist():
    sheet = Workbook().active
    sheet["A1"] = "value"
    assert isinstance(sheet._cells, dict)
    assert isinstance(sheet._cells[(1, 1)], Cell)
    assert sheet["A1"]._value == "value"
    assert hasattr(sheet["A1"], "_hyperlink")


def test_copy_worksheet_across_workbooks(tmp_path):
    source_workbook = Workbook()
    source = source_workbook.active
    source["A1"] = "Header"
    source["A1"].font = Font(bold=True)
    source["A1"].fill = PatternFill("solid", fgColor="FF0000")
    source["A1"].comment = Comment("note", "author")
    source["B2"] = 42
    source["C3"].hyperlink = "https://example.invalid"
    source.merge_cells("A4:C4")
    source.freeze_panes = "A2"
    source.column_dimensions["A"].width = 30
    validation = DataValidation(type="list", formula1='"Y,N"')
    source.add_data_validation(validation)
    validation.add("D2:D10")

    target_workbook = Workbook()
    copy_worksheet(target_workbook.active, source)
    target_workbook.save(tmp_path / "copy.xlsx")
    copied = load_workbook(tmp_path / "copy.xlsx").active
    assert copied["A1"].value == "Header" and copied["A1"].font.bold
    assert copied["A1"].fill.fgColor.rgb.endswith("FF0000")
    assert copied["A1"].comment.text == "note"
    assert copied["B2"].value == 42
    assert copied["C3"].hyperlink.target == "https://example.invalid"
    assert [str(cell_range) for cell_range in copied.merged_cells.ranges] == ["A4:C4"]
    assert copied.freeze_panes == "A2"
    assert copied.column_dimensions["A"].width == 30
    assert [str(item.sqref) for item in copied.data_validations.dataValidation] == ["D2:D10"]
