"""SWC 工具端到端测试：样例基线、大小写规则、错误阻止写出及各类新增对象。"""

from __future__ import annotations

import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from helpers import NS, ROOT, SWC, assert_matches_baseline, run_tool
from swc_fixtures import TWINS, run_prepared_twins, run_twins

def test_swc_sample(tmp_path):
    workbook = load_workbook(ROOT / "test/samples/swc/SwcTestPorts.xlsx")
    try:
        assert workbook.sheetnames == [
            "README", "Components", "Enumerations", "Arrays", "Structures",
            "SR-Interfaces", "CS-Interfaces", "P-Ports", "R-Ports", "Data Accesses",
        ]
        assert workbook["R-Ports"]["E1"].value == "Runnable Name"
        assert workbook["Data Accesses"]["B1"].value == "Runnable Name"
        for sheet_name in ("P-Ports", "R-Ports", "Data Accesses"):
            sheet = workbook[sheet_name]
            access_kind_column = next(
                cell.column_letter for cell in sheet[1] if cell.value == "Access Kind"
            )
            expected_range = f"{access_kind_column}2:{access_kind_column}1048576"
            assert any(
                expected_range in {str(cell_range) for cell_range in validation.sqref.ranges}
                for validation in sheet.data_validations.dataValidation
            )
    finally:
        workbook.close()

    result = run_tool(
        SWC, "--project-root", ROOT, "--config", ROOT / "test/samples/swc/TableToSwcArxml.toml",
        "--output-dir", tmp_path / "out", "--log-dir", tmp_path / "log",
    )
    assert result.returncode == 0, result.stderr
    produced = sorted((tmp_path / "out").rglob("*.arxml"))
    assert "Successfully wrote 2 ARXML file(s)" in result.stdout
    assert sorted(path.name for path in produced) == ["Horn.arxml", "LIT.arxml"]
    for path in produced:
        assert_matches_baseline(path, ROOT / "test/samples/swc/expected" / path.name)


def test_swc_warns_for_ignored_port_columns(tmp_path):
    workbook = load_workbook(ROOT / "test/samples/swc/SwcTestPorts.xlsx")
    for sheet_name, row_number, values in (
        (
            "P-Ports",
            2,
            {
                "Runnable Name": "UnusedRunnable",
                "Runnable Symbol": "unused_symbol",
                "Minimum Start Interval": 0.0,
                "Can Be Invoked Concurrently": "N",
            },
        ),
        (
            "R-Ports",
            2,
            {"Runnable Name": "UnusedRunnable", "Access Kind": "ReceiveByArgument"},
        ),
    ):
        sheet = workbook[sheet_name]
        columns = {cell.value: cell.column for cell in sheet[1]}
        for header, value in values.items():
            sheet.cell(row=row_number, column=columns[header], value=value)
    workbook.save(tmp_path / "ports.xlsx")
    workbook.close()

    result = run_tool(
        SWC, "--project-root", ROOT, "--config", ROOT / "test/samples/swc/TableToSwcArxml.toml",
        "--workbook", tmp_path / "ports.xlsx", "--output-dir", tmp_path / "out",
        "--log-dir", tmp_path / "log",
    )
    assert result.returncode == 0, result.stderr
    assert "do not apply to P-Port" in result.stderr
    assert "do not apply to R-Port" in result.stderr
    for path in (tmp_path / "out").rglob("*.arxml"):
        assert_matches_baseline(path, ROOT / "test/samples/swc/expected" / path.name)


def test_swc_accepts_case_insensitive_headers_and_sheet_names(tmp_path):
    workbook = load_workbook(ROOT / "test/samples/swc/SwcTestPorts.xlsx")
    final_titles = {}
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            if isinstance(cell.value, str) and sheet.title != "README":
                cell.value = cell.value.lower()
        if sheet.title != "README":
            final_titles[sheet] = sheet.title.lower()
    for sheet, title in final_titles.items():
        sheet.title = f"temporary-{sheet.title}"
        sheet.title = title
    workbook.save(tmp_path / "lowercase-ports.xlsx")
    workbook.close()

    result = run_tool(
        SWC, "--project-root", ROOT, "--config", ROOT / "test/samples/swc/TableToSwcArxml.toml",
        "--workbook", tmp_path / "lowercase-ports.xlsx", "--output-dir", tmp_path / "out",
        "--log-dir", tmp_path / "log",
    )
    assert result.returncode == 0, result.stderr
    for path in (tmp_path / "out").rglob("*.arxml"):
        assert_matches_baseline(path, ROOT / "test/samples/swc/expected" / path.name)


def component_block(data: bytes, name: str) -> bytes:
    start = data.index(f"<SHORT-NAME>{name}</SHORT-NAME>".encode())
    return data[start:data.index(b"</APPLICATION-SW-COMPONENT-TYPE>", start)]


def short_names(element: ET.Element, path: str) -> list[str]:
    return [child.findtext("ar:SHORT-NAME", namespaces=NS) for child in element.findall(path, NS)]


def test_swc_same_named_runnables_in_one_file(tmp_path):
    """两个组件的 runnable 同名时，访问点和服务器调用点只能写入工作簿指定的组件。

    数据元素与接口同名，同时覆盖同一文件内 SHORT-NAME 重复时不应误报多个定义。
    """

    source = tmp_path / "twins.arxml"
    source.write_text(TWINS, encoding="utf-8")
    (tmp_path / "arxml.lst").write_text("twins.arxml\n", encoding="utf-8")

    workbook = load_workbook(ROOT / "test/samples/swc/SwcTestPorts.xlsx")
    for sheet in workbook.worksheets[1:]:
        sheet.delete_rows(2, sheet.max_row)

    def append(sheet_name: str, **values: str) -> None:
        sheet = workbook[sheet_name]
        sheet.append(
            [
                values.get(str(cell.value).replace(" ", "_").replace("-", "_"))
                for cell in sheet[1]
            ]
        )

    append("Components", SW_C="A", Component_Path="/Pkg/A")
    append("Components", SW_C="B", Component_Path="/Pkg/B")
    append("P-Ports", SW_C="B", Port_Name="p_Out", Port_Interface_Ref="/Pkg/If",
           Create_RTE_Artifact="Y", Runnable_Name="Run")
    append("R-Ports", SW_C="B", Port_Name="r_Cs", Port_Interface_Ref="/Pkg/CsIf",
           Create_RTE_Artifact="Y", Runnable_Name="Run")
    workbook.save(tmp_path / "ports.xlsx")

    result = run_tool(
        SWC, "--project-root", tmp_path, "--arxml-file-list", tmp_path / "arxml.lst",
        "--workbook", tmp_path / "ports.xlsx", "--log-dir", tmp_path / "log",
        "--output-dir", tmp_path / "out",
    )
    assert result.returncode == 0, result.stderr

    output = (tmp_path / "out" / "twins.arxml").read_bytes()
    assert component_block(output, "A") == component_block(source.read_bytes(), "A")
    component = next(
        element for element in ET.fromstring(output).iter(f"{{{NS['ar']}}}APPLICATION-SW-COMPONENT-TYPE")
        if element.findtext("ar:SHORT-NAME", namespaces=NS) == "B")
    assert sorted(short_names(component, "ar:PORTS/*")) == ["p_Out", "r_Cs"]
    runnable = component.find("ar:INTERNAL-BEHAVIORS/*/ar:RUNNABLES/ar:RUNNABLE-ENTITY", NS)
    assert short_names(runnable, "ar:DATA-SEND-POINTS/*") == ["dsp_p_Out"]
    assert short_names(runnable, "ar:SERVER-CALL-POINTS/*") == ["scp_r_Cs_Op"]


def test_swc_port_access_kind_creates_implicit_sr_accesses(tmp_path):
    result = run_twins(tmp_path, {
        "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
        "P-Ports": [{"SW_C": "B", "Port_Name": "p_Write", "Port_Interface_Ref": "/Pkg/If",
                     "Create_RTE_Artifact": "Y", "Runnable_Name": "Run", "Access_Kind": "Write"}],
        "R-Ports": [{"SW_C": "B", "Port_Name": "r_Read", "Port_Interface_Ref": "/Pkg/If",
                     "Create_RTE_Artifact": "Y", "Runnable_Name": "Run", "Access_Kind": "Read"}],
    })
    assert result.returncode == 0, result.stderr

    output = ET.parse(tmp_path / "out" / "twins.arxml").getroot()
    component = next(
        element for element in output.iter(f"{{{NS['ar']}}}APPLICATION-SW-COMPONENT-TYPE")
        if element.findtext("ar:SHORT-NAME", namespaces=NS) == "B")
    runnable = component.find("ar:INTERNAL-BEHAVIORS/*/ar:RUNNABLES/ar:RUNNABLE-ENTITY", NS)
    assert short_names(runnable, "ar:DATA-WRITE-ACCESSS/*") == ["dwa_p_Write"]
    assert short_names(runnable, "ar:DATA-READ-ACCESSS/*") == ["dra_r_Read"]
    assert short_names(runnable, "ar:DATA-SEND-POINTS/*") == []


def test_swc_creates_missing_sr_access_container_in_schema_order(tmp_path):
    result = run_twins(tmp_path, {
        "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
        "R-Ports": [{"SW_C": "B", "Port_Name": "r_Read", "Port_Interface_Ref": "/Pkg/If",
                     "Create_RTE_Artifact": "Y", "Runnable_Name": "Run",
                     "Access_Kind": "ReceiveByValue"}],
    })
    assert result.returncode == 0, result.stderr

    output = ET.parse(tmp_path / "out" / "twins.arxml").getroot()
    runnable = next(
        element for element in output.iter(f"{{{NS['ar']}}}RUNNABLE-ENTITY")
        if element.findtext("ar:SHORT-NAME", namespaces=NS) == "Run"
        and element.find("ar:DATA-RECEIVE-POINT-BY-VALUES", NS) is not None
    )
    access_names = short_names(runnable, "ar:DATA-RECEIVE-POINT-BY-VALUES/*")
    assert access_names == ["drpval_r_Read"]
    child_tags = [child.tag.rsplit("}", maxsplit=1)[-1] for child in runnable]
    assert child_tags.index("DATA-RECEIVE-POINT-BY-VALUES") < child_tags.index("DATA-SEND-POINTS")


def test_swc_ignores_invalid_rte_fields_when_rte_artifacts_are_disabled(tmp_path):
    result = run_twins(tmp_path, {
        "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
        "P-Ports": [{
            "SW_C": "B",
            "Port_Name": "p_Test",
            "Port_Interface_Ref": "/Pkg/If",
            "Create_RTE_Artifact": "N",
            "Runnable_Name": "bad/name",
            "Runnable_Symbol": "not a symbol",
            "Access_Kind": "Unknown",
            "Minimum_Start_Interval": "not a number",
            "Can_Be_Invoked_Concurrently": "maybe",
        }],
    })
    assert result.returncode == 0, result.stderr
    assert "will be ignored" in result.stderr
    assert "unknown Access Kind" not in result.stderr
    assert "Minimum Start Interval must" not in result.stderr


def test_swc_checks_only_explicit_com_spec_fields_for_existing_port(tmp_path):
    result = run_twins(tmp_path, {
        "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
        "P-Ports": [{
            "SW_C": "B",
            "Port_Name": "p_Test",
            "Port_Interface_Ref": "/Pkg/If",
            "Transmission_Acknowledge_Timeout": "1",
            "Uses_End_To_End_Protection": "Y",
        }],
    })
    assert result.returncode == 0, result.stderr

    (tmp_path / "twins.arxml").write_bytes((tmp_path / "out" / "twins.arxml").read_bytes())
    workbook = load_workbook(tmp_path / "ports.xlsx")
    sheet = workbook["P-Ports"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(row=2, column=columns["Uses End To End Protection"]).value = None
    workbook.save(tmp_path / "ports.xlsx")
    workbook.close()
    workbook = load_workbook(tmp_path / "ports.xlsx")
    assert workbook["P-Ports"].cell(2, columns["Uses End To End Protection"]).value is None
    workbook.close()
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 0, result.stderr

    workbook = load_workbook(tmp_path / "ports.xlsx")
    sheet = workbook["P-Ports"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(row=2, column=columns["Transmission Acknowledge Timeout"]).value = None
    workbook.save(tmp_path / "ports.xlsx")
    workbook.close()
    workbook = load_workbook(tmp_path / "ports.xlsx")
    assert workbook["P-Ports"].cell(2, columns["Transmission Acknowledge Timeout"]).value is None
    workbook.close()
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 0, result.stderr

    workbook = load_workbook(tmp_path / "ports.xlsx")
    sheet = workbook["P-Ports"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(row=2, column=columns["Transmission Acknowledge Timeout"], value=2)
    workbook.save(tmp_path / "ports.xlsx")
    workbook.close()
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 1
    assert "Transmission Acknowledge Timeout" in result.stderr


@pytest.mark.parametrize("field,value,expected", [
    ("Transmission Acknowledge Timeout", 1, 0),
    ("Transmission Acknowledge Timeout", 0, 1),
    ("Uses End To End Protection", "Y", 0),
    ("Uses End To End Protection", "N", 1),
])
def test_existing_sender_com_spec_checks_only_selected_field(tmp_path, field, value, expected):
    result = run_twins(tmp_path, {
        "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
        "P-Ports": [{
            "SW_C": "B", "Port_Name": "p_Test", "Port_Interface_Ref": "/Pkg/If",
            "Transmission_Acknowledge_Timeout": "1", "Uses_End_To_End_Protection": "Y",
        }],
    })
    assert result.returncode == 0, result.stderr
    (tmp_path / "twins.arxml").write_bytes((tmp_path / "out" / "twins.arxml").read_bytes())
    workbook = load_workbook(tmp_path / "ports.xlsx")
    sheet = workbook["P-Ports"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for header in ("Transmission Acknowledge Timeout", "Uses End To End Protection"):
        sheet.cell(2, columns[header]).value = value if header == field else None
    workbook.save(tmp_path / "ports.xlsx")
    workbook.close()
    result = run_prepared_twins(tmp_path)
    assert result.returncode == expected, result.stderr
    if expected:
        assert field in result.stderr


def test_swc_port_access_kind_must_match_direction(tmp_path):
    result = run_twins(tmp_path, {
        "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
        "P-Ports": [{"SW_C": "B", "Port_Name": "p_Bad", "Port_Interface_Ref": "/Pkg/If",
                     "Create_RTE_Artifact": "Y", "Runnable_Name": "Run", "Access_Kind": "Read"}],
    })
    assert result.returncode == 1
    assert "Access Kind Read requires a R-Port" in result.stderr
    assert not (tmp_path / "out").exists()


def test_swc_only_required_sheets_and_reordered_columns(tmp_path):
    workbook = load_workbook(ROOT / "test/samples/swc/SwcTestPorts.xlsx")
    for sheet in list(workbook):
        if sheet.title not in ("Components", "P-Ports", "R-Ports"):
            workbook.remove(sheet)
    for sheet in workbook:
        rows = [tuple(reversed(row)) for row in sheet.iter_rows(values_only=True)]
        sheet.delete_rows(1, sheet.max_row)
        for row in rows:
            sheet.append(row)
        # These test inputs exercise the parser, not Excel's table editor.
        sheet.tables.clear()
    workbook.save(tmp_path / "minimal.xlsx")
    workbook.close()
    result = run_tool(
        SWC, "--project-root", ROOT, "--config", ROOT / "test/samples/swc/TableToSwcArxml.toml",
        "--workbook", tmp_path / "minimal.xlsx", "--output-dir", tmp_path / "out",
        "--log-dir", tmp_path / "log",
    )
    assert result.returncode == 0, result.stderr
    produced = sorted((tmp_path / "out").rglob("*.arxml"))
    assert [path.name for path in produced] == ["LIT.arxml", "Horn.arxml"]
    for path in produced:
        assert_matches_baseline(path, ROOT / "test/samples/swc/expected" / path.name)


@pytest.mark.parametrize("sheet", ["P-Ports", "Data Accesses"])
def test_swc_missing_components_still_fails(tmp_path, sheet):
    result = run_twins(tmp_path, {
        sheet: [{"SW_C": "B", "Port_Name": "p_Test", "Port_Interface_Ref": "/Pkg/If",
                 "Runnable_Name": "Run"}],
    }, minimal=True)
    assert result.returncode == 1
    assert "SW-C B is not defined in the Components sheet" in result.stderr
    assert not (tmp_path / "out").exists()


@pytest.mark.parametrize("create_ports", [False, True])
def test_swc_zero_argument_operation(tmp_path, create_ports):
    rows = {
        "CS-Interfaces": [{"Interface_Path": "/Pkg/NewCs", "Operation_Name": "Reset",
                           "Target_ARXML": "twins.arxml"}],
    }
    if create_ports:
        rows.update({
            "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
            "P-Ports": [{"SW_C": "B", "Port_Name": "p_Reset", "Port_Interface_Ref": "/Pkg/NewCs",
                         "Create_RTE_Artifact": "Y"}],
            "R-Ports": [{"SW_C": "B", "Port_Name": "r_Reset", "Port_Interface_Ref": "/Pkg/NewCs",
                         "Create_RTE_Artifact": "Y", "Runnable_Name": "Run"}],
        })
    result = run_twins(tmp_path, rows, minimal=True)
    assert result.returncode == 0, result.stderr
    root = ET.parse(tmp_path / "out" / "twins.arxml").getroot()
    interface = next(element for element in root.findall(".//ar:CLIENT-SERVER-INTERFACE", NS)
                     if element.findtext("ar:SHORT-NAME", namespaces=NS) == "NewCs")
    assert short_names(interface, "ar:OPERATIONS/*") == ["Reset"]
    assert interface.findall(".//ar:ARGUMENT-DATA-PROTOTYPE", NS) == []
    if create_ports:
        assert "RE_p_Reset_Reset" in short_names(root, ".//ar:RUNNABLE-ENTITY")
        assert "oie_p_Reset_Reset" in short_names(root, ".//ar:OPERATION-INVOKED-EVENT")
        assert "scp_r_Reset_Reset" in short_names(root, ".//ar:SYNCHRONOUS-SERVER-CALL-POINT")


@pytest.mark.parametrize("argument", [
    {"Argument_Name": "value"},
    {"Direction": "IN"},
    {"Data_Type_Ref": "/Pkg/uint8"},
    {"Argument_Name": "value", "Direction": "IN"},
    {"Argument_Name": "value", "Data_Type_Ref": "/Pkg/uint8"},
    {"Direction": "IN", "Data_Type_Ref": "/Pkg/uint8"},
])
def test_swc_partial_cs_argument_blocks_writes(tmp_path, argument):
    result = run_twins(tmp_path, {
        "CS-Interfaces": [{"Interface_Path": "/Pkg/NewCs", "Operation_Name": "Reset",
                           "Target_ARXML": "twins.arxml", **argument}],
    }, minimal=True)
    assert result.returncode == 1
    assert "must be all filled or all empty" in result.stderr
    assert not (tmp_path / "out").exists()


@pytest.mark.parametrize("order", ["empty-first", "argument-first", "empty-twice"])
def test_swc_conflicting_zero_argument_declarations(tmp_path, order):
    empty = {"Interface_Path": "/Pkg/NewCs", "Operation_Name": "Reset", "Target_ARXML": "twins.arxml"}
    argument = {**empty, "Argument_Name": "value", "Direction": "IN", "Data_Type_Ref": "/Pkg/uint8"}
    rows = [empty, argument] if order == "empty-first" else (
        [argument, empty] if order == "argument-first" else [empty, empty]
    )
    result = run_twins(tmp_path, {"CS-Interfaces": rows}, minimal=True)
    assert result.returncode == 1
    assert "duplicate or conflicting zero-argument declaration" in result.stderr
    assert not (tmp_path / "out").exists()


def test_swc_invalid_rte_field_is_ignored_when_rte_artifacts_are_disabled(tmp_path):
    result = run_twins(tmp_path, {
        "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
        "P-Ports": [{"SW_C": "B", "Port_Name": "p_Test", "Port_Interface_Ref": "/Pkg/If",
                     "Create_RTE_Artifact": "N", "Minimum_Start_Interval": "abc"}],
    }, minimal=True)
    assert result.returncode == 0, result.stderr
    assert "Minimum Start Interval do not apply" in result.stderr
    assert (tmp_path / "out" / "twins.arxml").exists()


@pytest.mark.parametrize("timeout,e2e,expected", [(None, "N", False), ("0", "N", True), (None, "Y", True)])
def test_swc_comspec_trigger_values(tmp_path, timeout, e2e, expected):
    result = run_twins(tmp_path, {
        "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
        "P-Ports": [{"SW_C": "B", "Port_Name": "p_Test", "Port_Interface_Ref": "/Pkg/If",
                     "Transmission_Acknowledge_Timeout": timeout, "Uses_End_To_End_Protection": e2e}],
    }, minimal=True)
    assert result.returncode == 0, result.stderr
    root = ET.parse(tmp_path / "out" / "twins.arxml").getroot()
    assert bool(root.findall(".//ar:NONQUEUED-SENDER-COM-SPEC", NS)) == expected