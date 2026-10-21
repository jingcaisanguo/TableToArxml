"""Regression tests for global definitions, output protection and invalid names."""

import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from helpers import NS, SWC, run_tool
from swc_fixtures import run_prepared_twins, run_twins, write_twins

PORTS = {
    "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
    "P-Ports": [{"SW_C": "B", "Port_Name": "p_Test", "Port_Interface_Ref": "/Pkg/If"}],
}
EMPTY_PACKAGE = (
    '<AUTOSAR xmlns="http://autosar.org/schema/r4.0"><AR-PACKAGES>'
    "<AR-PACKAGE><SHORT-NAME>Pkg</SHORT-NAME><ELEMENTS></ELEMENTS>"
    "</AR-PACKAGE></AR-PACKAGES></AUTOSAR>"
)
DUPLICATE_SR_INTERFACE = """<AUTOSAR xmlns="http://autosar.org/schema/r4.0">
  <AR-PACKAGES><AR-PACKAGE><SHORT-NAME>Pkg</SHORT-NAME><ELEMENTS>
    <SENDER-RECEIVER-INTERFACE><SHORT-NAME>If</SHORT-NAME><IS-SERVICE>false</IS-SERVICE>
      <DATA-ELEMENTS><VARIABLE-DATA-PROTOTYPE><SHORT-NAME>If</SHORT-NAME>
        <TYPE-TREF DEST="IMPLEMENTATION-DATA-TYPE">/Pkg/uint8</TYPE-TREF>
      </VARIABLE-DATA-PROTOTYPE></DATA-ELEMENTS>
    </SENDER-RECEIVER-INTERFACE>
  </ELEMENTS></AR-PACKAGE></AR-PACKAGES>
</AUTOSAR>"""


def test_broken_candidate_prevents_uniqueness_claim(tmp_path):
    write_twins(tmp_path, PORTS, minimal=True)
    (tmp_path / "broken.arxml").write_text(
        "<AUTOSAR><AR-PACKAGES><AR-PACKAGE><SHORT-NAME>Pkg</SHORT-NAME>"
        "<ELEMENTS><APPLICATION-SW-COMPONENT-TYPE><SHORT-NAME>B</SHORT-NAME>",
        encoding="utf-8",
    )
    (tmp_path / "arxml.lst").write_text("twins.arxml\nbroken.arxml\n", encoding="utf-8")
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 1
    assert "parser failures" in result.stderr
    assert "broken.arxml" in result.stderr
    assert not (tmp_path / "out").exists()


def test_non_utf8_input_is_not_silently_ignored(tmp_path):
    write_twins(tmp_path, PORTS, minimal=True)
    (tmp_path / "broken.arxml").write_bytes(b"\xff\xff")
    (tmp_path / "arxml.lst").write_text("twins.arxml\nbroken.arxml\n", encoding="utf-8")
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 1
    assert "Cannot read UTF-8 ARXML" in result.stderr
    assert not (tmp_path / "out").exists()


def test_interface_cannot_duplicate_a_definition_in_another_file(tmp_path):
    write_twins(tmp_path, {
        "SR-Interfaces": [{"Interface_Path": "/Pkg/If", "Data_Element_Name": "If",
                           "Data_Type_Ref": "/Pkg/Other", "Is_Service": "Y",
                           "Target_ARXML": "second.arxml"}],
    }, minimal=True)
    (tmp_path / "second.arxml").write_text(EMPTY_PACKAGE, encoding="utf-8")
    (tmp_path / "arxml.lst").write_text("twins.arxml\nsecond.arxml\n", encoding="utf-8")
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 1
    assert "SR interface /Pkg/If in" in result.stderr
    assert "already exists with a different definition" in result.stderr
    assert not (tmp_path / "out").exists()


def test_identical_interface_definitions_in_multiple_files_are_accepted(tmp_path):
    write_twins(tmp_path, PORTS, minimal=True)
    (tmp_path / "second.arxml").write_text(DUPLICATE_SR_INTERFACE, encoding="utf-8")
    (tmp_path / "arxml.lst").write_text("twins.arxml\nsecond.arxml\n", encoding="utf-8")

    result = run_prepared_twins(tmp_path)

    assert result.returncode == 0, result.stderr
    assert "Multiple definitions found for interface /Pkg/If" in result.stderr
    assert "using" in result.stderr
    assert (tmp_path / "out" / "twins.arxml").exists()


def test_existing_interface_target_is_reused_despite_duplicate_in_another_file(tmp_path):
    write_twins(tmp_path, {
        "SR-Interfaces": [{
            "Interface_Path": "/Pkg/If",
            "Data_Element_Name": "If",
            "Data_Type_Ref": "/Pkg/uint8",
            "Target_ARXML": "twins.arxml",
        }],
    }, minimal=True)
    (tmp_path / "second.arxml").write_text(DUPLICATE_SR_INTERFACE, encoding="utf-8")
    (tmp_path / "arxml.lst").write_text("twins.arxml\nsecond.arxml\n", encoding="utf-8")

    result = run_prepared_twins(tmp_path)

    assert result.returncode == 0, result.stderr
    assert "already exists with the same definition" in result.stderr


def test_new_type_cannot_collide_with_an_interface_in_another_file(tmp_path):
    write_twins(tmp_path, {
        "Arrays": [{"Type_Path": "/Pkg/If", "Element_Name": "Element",
                    "Element_Type_Ref": "/Pkg/uint8", "Array_Size": "2",
                    "Target_ARXML": "second.arxml"}],
    }, minimal=True)
    (tmp_path / "second.arxml").write_text(EMPTY_PACKAGE, encoding="utf-8")
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 1
    assert "array /Pkg/If" in result.stderr
    assert "already exists with a different definition" in result.stderr
    assert not (tmp_path / "out").exists()


def test_target_files_outside_list_participate_in_conflict_checks(tmp_path):
    write_twins(tmp_path, {
        "SR-Interfaces": [{"Interface_Path": "/Pkg/NewIf", "Data_Element_Name": "value",
                           "Data_Type_Ref": "/Pkg/uint8", "Target_ARXML": "second.arxml"}],
        "CS-Interfaces": [{"Interface_Path": "/Pkg/NewIf", "Operation_Name": "Op",
                           "Target_ARXML": "third.arxml"}],
    }, minimal=True)
    for name in ("second.arxml", "third.arxml"):
        (tmp_path / name).write_text(EMPTY_PACKAGE, encoding="utf-8")
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 1
    assert "CS interface /Pkg/NewIf" in result.stderr
    assert "already exists with a different definition" in result.stderr
    assert not (tmp_path / "out").exists()


def test_duplicate_component_uses_first_arxml_file(tmp_path):
    write_twins(tmp_path, {
        "Components": [{"SW_C": "B", "Component_Path": "/Pkg/B"}],
        "P-Ports": [{"SW_C": "B", "Port_Name": "p_First", "Port_Interface_Ref": "/Pkg/If"}],
    }, minimal=True)
    (tmp_path / "second.arxml").write_bytes((tmp_path / "twins.arxml").read_bytes())
    (tmp_path / "arxml.lst").write_text("second.arxml\ntwins.arxml\n", encoding="utf-8")

    result = run_prepared_twins(tmp_path)

    assert result.returncode == 0, result.stderr
    assert "Multiple definitions found for component /Pkg/B" in result.stderr
    first_output = tmp_path / "out" / "second.arxml"
    assert first_output.exists()
    assert b"<SHORT-NAME>p_First</SHORT-NAME>" in first_output.read_bytes()
    assert not (tmp_path / "out" / "twins.arxml").exists()


@pytest.mark.parametrize("sheet,record", [
    ("Arrays", {
        "Type_Path": "/Pkg/NewArray", "Element_Name": "item",
        "Element_Type_Ref": "/Pkg/uint8", "Array_Size": "2",
    }),
    ("Structures", {
        "Type_Path": "/Pkg/NewStructure", "Member_Name": "item",
        "Member_Type_Ref": "/Pkg/uint8",
    }),
    ("Enumerations", {
        "Type_Path": "/Pkg/NewEnum",
        "Implementation_Data_Type_Ref": "/Pkg/uint8",
        "Value": "0",
        "Label": "Off",
        "Compu_Method_Ref": "/Pkg/NewEnumCompu",
        "Data_Constr_Ref": "/Pkg/NewEnumConstr",
    }),
])
def test_existing_type_in_first_arxml_is_reused_instead_of_duplicated(
    tmp_path, sheet, record
):
    initial = {key: value for key, value in record.items()}
    initial["Target_ARXML"] = "twins.arxml"
    result = run_twins(tmp_path, {sheet: [initial]}, minimal=True)
    assert result.returncode == 0, result.stderr

    generated = (tmp_path / "out" / "twins.arxml").read_bytes()
    (tmp_path / "twins.arxml").write_bytes(generated)
    (tmp_path / "second.arxml").write_bytes(generated)
    (tmp_path / "target.arxml").write_text(EMPTY_PACKAGE, encoding="utf-8")
    (tmp_path / "arxml.lst").write_text("twins.arxml\nsecond.arxml\n", encoding="utf-8")

    workbook = load_workbook(tmp_path / "ports.xlsx")
    sheet_data = workbook[sheet]
    columns = {cell.value: cell.column for cell in sheet_data[1]}
    sheet_data.cell(row=2, column=columns["Target ARXML"], value="target.arxml")
    workbook.save(tmp_path / "ports.xlsx")
    workbook.close()
    result = run_prepared_twins(tmp_path)

    assert result.returncode == 0, result.stderr
    assert "already exists with the same definition" in result.stderr
    assert not (tmp_path / "out" / "target.arxml").exists()


def test_output_directory_cannot_implicitly_enable_in_place(tmp_path):
    write_twins(tmp_path, PORTS, minimal=True)
    original = (tmp_path / "twins.arxml").read_bytes()
    result = run_tool(
        SWC, "--project-root", tmp_path, "--arxml-file-list", tmp_path / "arxml.lst",
        "--workbook", tmp_path / "ports.xlsx", "--log-dir", tmp_path / "log",
        "--output-dir", tmp_path,
    )
    assert result.returncode == 1
    assert "use --in-place explicitly" in result.stderr
    assert (tmp_path / "twins.arxml").read_bytes() == original


def test_blank_header_blocks_output_instead_of_disabling_artifacts(tmp_path):
    write_twins(tmp_path, PORTS, minimal=True)
    path = tmp_path / "ports.xlsx"
    workbook = load_workbook(path)
    sheet = workbook["P-Ports"]
    sheet.insert_cols(4)
    sheet["E2"] = "Y"
    workbook.save(path)
    workbook.close()
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 1
    assert "Skipped worksheet P-Ports" in result.stderr
    assert not (tmp_path / "out").exists()


@pytest.mark.parametrize("title", ["Data-Accesses", "Custom Data"])
def test_unknown_worksheet_blocks_output_even_with_valid_ports(tmp_path, title):
    write_twins(tmp_path, PORTS, minimal=True)
    path = tmp_path / "ports.xlsx"
    workbook = load_workbook(path)
    workbook.create_sheet(title).append(["SW-C", "Runnable Name", "Port Name"])
    workbook.save(path)
    workbook.close()
    result = run_prepared_twins(tmp_path)
    assert result.returncode == 1
    assert f"Unknown worksheet {title}" in result.stderr
    assert not (tmp_path / "out").exists()


@pytest.mark.parametrize("name", ["bad/name", "1port", "P" * 129])
def test_invalid_short_name_never_reaches_output(tmp_path, name):
    result = run_twins(tmp_path, {
        **PORTS, "P-Ports": [{"SW_C": "B", "Port_Name": name, "Port_Interface_Ref": "/Pkg/If"}],
    }, minimal=True)
    assert result.returncode == 1
    assert "AUTOSAR SHORT-NAME" in result.stderr
    assert not (tmp_path / "out").exists()


def test_long_generated_cs_names_are_bounded_and_idempotent(tmp_path):
    name = "P" * 128
    result = run_twins(tmp_path, {
        "Components": PORTS["Components"],
        "P-Ports": [{"SW_C": "B", "Port_Name": name, "Port_Interface_Ref": "/Pkg/CsIf",
                     "Create_RTE_Artifact": "Y"}],
    }, minimal=True)
    assert result.returncode == 0, result.stderr
    produced = tmp_path / "out" / "twins.arxml"
    tree = ET.parse(produced)
    assert all(len(node.text) <= 128 for node in tree.findall(".//ar:SHORT-NAME", NS))
    (tmp_path / "twins.arxml").write_bytes(produced.read_bytes())
    repeated = run_prepared_twins(tmp_path)
    assert repeated.returncode == 0, repeated.stderr
    assert "already exists with the same definition" in repeated.stderr
