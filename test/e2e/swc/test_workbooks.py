"""工作簿产物的布局和再生成回归，不依赖 Excel GUI。"""

import subprocess
import sys
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from helpers import ROOT, SWC, run_tool

SAMPLE = ROOT / "test" / "tableToSwcArxml" / "SwcTestPorts.xlsx"
TEMPLATE = ROOT / "docs" / "user" / "tableToSwcArxml" / "SwcTemplate.xlsx"
SHEETS = [
    "README", "Components", "Enumerations", "Arrays", "Structures",
    "SR-Interfaces", "CS-Interfaces", "P-Ports", "R-Ports", "Data Accesses",
]
COMMON_PORT_HEADERS = [
    "SW-C", "Port Name", "Port Interface Ref", "Create RTE Artifact", "Runnable Name", "Access Kind",
]
BOOLEAN_CHOICES = {"Y", "N", "TRUE", "FALSE", "YES", "NO", "1", "0"}


@pytest.mark.parametrize("path", [SAMPLE, TEMPLATE], ids=["sample", "blank"])
def test_workbook_layout_and_validation(path):
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == SHEETS
        assert workbook.active.title == "README"
        readme = workbook["README"]
        assert readme.freeze_panes is None
        assert readme.auto_filter.ref is None
        assert readme.sheet_view.topLeftCell == "A1"
        assert readme.sheet_view.selection[0].activeCell == "A1"
        assert all(readme.row_dimensions[row].height >= 30 for row in range(1, readme.max_row + 1))
        sections = {
            row[0].value: row[0].row for row in readme
            if row[0].hyperlink and row[0].value in SHEETS[1:]
        }
        assert list(sections) == SHEETS[1:]
        assert workbook.defined_names["SwcComponentKeys"].attr_text == "SwcComponents[SW-C]"
        for name in SHEETS[1:]:
            sheet = workbook[name]
            assert sheet.sheet_view.showGridLines is False
            assert sheet.sheet_view.topLeftCell == "A1"
            assert all(selection.activeCell == "A2" for selection in sheet.sheet_view.selection)
            assert sheet.freeze_panes == ("C2" if name in ("P-Ports", "R-Ports", "Data Accesses") else "A2")
            assert sheet["A1"].hyperlink.location == f"'README'!A{sections[name]}"
            assert sheet.sheet_properties.tabColor is not None
            assert sheet.auto_filter.ref is None
            assert len(sheet.tables) == 1
            table = next(iter(sheet.tables.values()))
            assert table.ref == sheet.dimensions
            assert table.autoFilter.ref == table.ref
            assert [column.name for column in table.tableColumns] == [cell.value for cell in sheet[1]]
            assert sheet.row_dimensions[1].height >= 48
            assert sheet.print_title_rows == "$1:$1"
            assert sheet.page_setup.orientation == "landscape"
            assert sheet.page_setup.fitToWidth == 1
            for cell in sheet[1]:
                assert cell.font.bold
                assert cell.fill.fgColor.rgb in ("001F4E78", "009C6500", "00516272")
                assert cell.alignment.wrap_text
                assert cell.comment and "默认" in cell.comment.text
                assert 12 <= sheet.column_dimensions[cell.column_letter].width <= 60
            for validation in sheet.data_validations.dataValidation:
                assert validation.showInputMessage and validation.showErrorMessage
                assert validation.errorStyle == "stop"
                assert validation.allow_blank
                assert len(validation.prompt) <= 255
                for cell_range in validation.sqref.ranges:
                    first_col, first_row, last_col, last_row = range_boundaries(str(cell_range))
                    assert first_col == last_col
                    assert first_row == 2 and last_row == 1048576
                if validation.formula1.startswith('"Y,N,'):
                    assert set(validation.formula1.strip('"').split(",")) == BOOLEAN_CHOICES
            if path == TEMPLATE:
                assert sheet.max_row == 2
                assert not any(cell.value is not None for row in sheet.iter_rows(min_row=2) for cell in row)
        for name in ("P-Ports", "R-Ports"):
            sheet = workbook[name]
            assert [cell.value for cell in sheet[1]][:6] == COMMON_PORT_HEADERS
            assert any(v.formula1 == "=SwcComponentKeys" for v in sheet.data_validations.dataValidation)
            assert any(str(v.sqref) == "F2:F1048576" for v in sheet.data_validations.dataValidation)
        assert [workbook["P-Ports"].cell(1, col).value for col in range(7, 12)] == [
            "Transmission Acknowledge Timeout", "Uses End To End Protection",
            "Runnable Symbol", "Minimum Start Interval", "Can Be Invoked Concurrently",
        ]
        assert workbook["Enumerations"]["D1"].value == "Implementation Data Type Ref"
        assert "不是 SW-BASE-TYPE" in workbook["Enumerations"]["D1"].comment.text
        assert workbook["P-Ports"].column_dimensions["C"].width == workbook["R-Ports"].column_dimensions["C"].width
        for name in ("Enumerations", "Arrays", "P-Ports", "R-Ports"):
            assert any(v.type == "custom" for v in workbook[name].data_validations.dataValidation)
    finally:
        workbook.close()


def test_workbooks_are_reproducible(tmp_path):
    result = subprocess.run(
        [sys.executable, str(ROOT / "tools" / "generate_swc_workbooks.py"), "--output-dir", str(tmp_path)],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    )
    assert result.returncode == 0, result.stderr
    for source in (SAMPLE, TEMPLATE):
        with ZipFile(source) as expected, ZipFile(tmp_path / source.name) as actual:
            names = [name for name in expected.namelist() if name.startswith("xl/")]
            assert names == [name for name in actual.namelist() if name.startswith("xl/")]
            for name in names:
                assert expected.read(name) == actual.read(name), (source.name, name)


def test_blank_template_has_no_implicit_operations(tmp_path):
    result = run_tool(
        SWC, "--config", SAMPLE.parent / "TableToSwcArxml.json",
        "--workbook", TEMPLATE, "--output-dir", tmp_path / "out", "--log-dir", tmp_path / "log",
    )
    assert result.returncode == 1
    assert "No usable records found in workbook" in result.stderr
    assert not (tmp_path / "out").exists()
