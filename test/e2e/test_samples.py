"""端到端测试：运行两个工具的样例并与 expected/ 下的基线比对。"""

from __future__ import annotations

import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SYSTEM = ROOT / "table_to_system_arxml" / "table_to_system_arxml.py"
SWC = ROOT / "table_to_swc_arxml" / "table_to_swc_arxml.py"
NS = {"ar": "http://autosar.org/schema/r4.0"}


def run_tool(script: Path, *args: object) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        cwd=ROOT, capture_output=True, text=True, encoding="utf-8", errors="replace",
    )


def cells(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}
    finally:
        workbook.close()


def lf_bytes(path: Path) -> bytes:
    # 基线和检出的样例可能是 CRLF 或 LF，比较前统一换行符。
    return path.read_bytes().replace(b"\r\n", b"\n")


def assert_matches_baseline(actual: Path, expected: Path) -> None:
    assert actual.is_file(), f"missing output {actual}"
    if expected.suffix == ".xlsx":
        assert cells(actual) == cells(expected), f"{actual.name} differs from {expected}"
    else:
        assert lf_bytes(actual) == lf_bytes(expected), f"{actual.name} differs from {expected}"


@pytest.mark.parametrize("variant, prefix", [("BcmTest", "Bcm"), ("McTest", "Mc")])
def test_system_sample(tmp_path, variant, prefix):
    result = run_tool(
        SYSTEM, "--config", ROOT / f"test/tableToSystemArxml/{prefix}TableToSystemArxmlConfig.json",
        "--output-dir", tmp_path / "out", "--log-dir", tmp_path / "log",
    )
    assert result.returncode == 0, result.stderr
    expected_dir = ROOT / "test/tableToSystemArxml/expected" / variant
    assert sorted(path.name for path in (tmp_path / "out").iterdir()) == sorted(
        path.name for path in expected_dir.iterdir())
    for expected in expected_dir.iterdir():
        assert_matches_baseline(tmp_path / "out" / expected.name, expected)


def test_swc_sample(tmp_path):
    result = run_tool(
        SWC, "--project-root", ROOT, "--config", ROOT / "test/tableToSwcArxml/TableToSwcArxml.json",
        "--output-dir", tmp_path / "out", "--log-dir", tmp_path / "log",
    )
    assert result.returncode == 0, result.stderr
    produced = sorted((tmp_path / "out").rglob("*.arxml"))
    assert sorted(path.name for path in produced) == ["Horn.arxml", "LIT.arxml"]
    for path in produced:
        assert_matches_baseline(path, ROOT / "test/tableToSwcArxml/expected" / path.name)


COMPONENT = """        <APPLICATION-SW-COMPONENT-TYPE>
          <SHORT-NAME>{name}</SHORT-NAME>
          <PORTS></PORTS>
          <INTERNAL-BEHAVIORS>
            <SWC-INTERNAL-BEHAVIOR>
              <SHORT-NAME>IB</SHORT-NAME>
              <RUNNABLES>
                <RUNNABLE-ENTITY>
                  <SHORT-NAME>Run</SHORT-NAME>
                  <MINIMUM-START-INTERVAL>0.0</MINIMUM-START-INTERVAL>
                  <CAN-BE-INVOKED-CONCURRENTLY>false</CAN-BE-INVOKED-CONCURRENTLY>
                  <DATA-SEND-POINTS></DATA-SEND-POINTS>
                  <SYMBOL>{name}_Run</SYMBOL>
                </RUNNABLE-ENTITY>
              </RUNNABLES>
            </SWC-INTERNAL-BEHAVIOR>
          </INTERNAL-BEHAVIORS>
        </APPLICATION-SW-COMPONENT-TYPE>
"""

TWINS = """<?xml version="1.0" encoding="UTF-8"?>
<AUTOSAR xmlns="http://autosar.org/schema/r4.0" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:schemaLocation="http://autosar.org/schema/r4.0 AUTOSAR_4-3-0.xsd">
  <AR-PACKAGES>
    <AR-PACKAGE>
      <SHORT-NAME>Pkg</SHORT-NAME>
      <ELEMENTS>
        <SENDER-RECEIVER-INTERFACE>
          <SHORT-NAME>If</SHORT-NAME>
          <IS-SERVICE>false</IS-SERVICE>
          <DATA-ELEMENTS>
            <VARIABLE-DATA-PROTOTYPE>
              <SHORT-NAME>de</SHORT-NAME>
              <TYPE-TREF DEST="IMPLEMENTATION-DATA-TYPE">/Pkg/uint8</TYPE-TREF>
            </VARIABLE-DATA-PROTOTYPE>
          </DATA-ELEMENTS>
        </SENDER-RECEIVER-INTERFACE>
        <CLIENT-SERVER-INTERFACE>
          <SHORT-NAME>CsIf</SHORT-NAME>
          <IS-SERVICE>false</IS-SERVICE>
          <OPERATIONS>
            <CLIENT-SERVER-OPERATION>
              <SHORT-NAME>Op</SHORT-NAME>
            </CLIENT-SERVER-OPERATION>
          </OPERATIONS>
        </CLIENT-SERVER-INTERFACE>
""" + COMPONENT.format(name="A") + COMPONENT.format(name="B") + """      </ELEMENTS>
    </AR-PACKAGE>
  </AR-PACKAGES>
</AUTOSAR>
"""


def component_block(data: bytes, name: str) -> bytes:
    start = data.index(f"<SHORT-NAME>{name}</SHORT-NAME>".encode())
    return data[start:data.index(b"</APPLICATION-SW-COMPONENT-TYPE>", start)]


def short_names(element: ET.Element, path: str) -> list[str]:
    return [child.findtext("ar:SHORT-NAME", namespaces=NS) for child in element.findall(path, NS)]


def test_swc_same_named_runnables_in_one_file(tmp_path):
    """两个组件的 runnable 同名时，访问点和服务器调用点只能写入工作簿指定的组件。"""

    source = tmp_path / "twins.arxml"
    source.write_text(TWINS, encoding="utf-8")
    (tmp_path / "arxml.lst").write_text("twins.arxml\n", encoding="utf-8")

    workbook = load_workbook(ROOT / "test/tableToSwcArxml/SwcTestPorts.xlsx")
    for sheet in workbook.worksheets[1:]:
        sheet.delete_rows(2, sheet.max_row)

    def append(sheet_name: str, **values: str) -> None:
        sheet = workbook[sheet_name]
        sheet.append([values.get(str(cell.value).replace(" ", "_")) for cell in sheet[1]])

    append("Components", SWC="A", Component_Path="/Pkg/A")
    append("Components", SWC="B", Component_Path="/Pkg/B")
    append("P-Ports", SWC="B", Port_Name="p_Out", Port_Interface_Ref="/Pkg/If",
           Create_RTE_Artifact="Y", Runnable_Name="Run")
    append("R-Ports", SWC="B", Port_Name="r_Cs", Port_Interface_Ref="/Pkg/CsIf",
           Create_RTE_Artifact="Y", Runnable_Ref="Run")
    workbook.save(tmp_path / "ports.xlsx")

    result = run_tool(
        SWC, "--project-root", tmp_path, "--arxml-file-list", tmp_path / "arxml.lst",
        "--workbook", tmp_path / "ports.xlsx", "--log-dir", tmp_path / "log",
        "--output-dir", tmp_path / "out",
    )
    assert result.returncode == 0, result.stderr

    output = (tmp_path / "out" / "twins.arxml").read_bytes()
    assert component_block(output, "A") == component_block(source.read_bytes(), "A")
    component = next(
        element for element in ET.fromstring(output).iter(f"{{{NS['ar']}}}APPLICATION-SW-COMPONENT-TYPE")
        if element.findtext("ar:SHORT-NAME", namespaces=NS) == "B")
    assert sorted(short_names(component, "ar:PORTS/*")) == ["p_Out", "r_Cs"]
    runnable = component.find("ar:INTERNAL-BEHAVIORS/*/ar:RUNNABLES/ar:RUNNABLE-ENTITY", NS)
    assert short_names(runnable, "ar:DATA-SEND-POINTS/*") == ["dsp_p_Out"]
    assert short_names(runnable, "ar:SERVER-CALL-POINTS/*") == ["scp_r_Cs_Op"]
