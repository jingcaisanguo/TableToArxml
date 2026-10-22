"""System 输出的 AUTOSAR 浅校验：多重性、不可拆分聚合、命名空间、引用元类型与版本提示。"""

from __future__ import annotations

import json
import xml.etree.ElementTree as ET

import pytest
from openpyxl import Workbook

from helpers import NS, SYSTEM, run_tool
from system_fixtures import write_config, write_model, write_white_table

AR = NS["ar"]
REFERENCES = [["A", "A", "/Pkg/A", "Reference"], ["B", "B", "/Pkg/B", "Reference"]]


def edit_model(tmp_path, *fragments: str, system_children: str = "", composition_children: str = "") -> None:
    """在 /Pkg 下追加元素，并可向 /Pkg/Sys、/Pkg/Top 追加子元素。"""

    write_model(tmp_path)
    path = tmp_path / "model.arxml"
    text = path.read_text(encoding="utf-8")
    elements = "\n".join(["<ECU-INSTANCE><SHORT-NAME>Ecu</SHORT-NAME></ECU-INSTANCE>", *fragments])
    text = text.replace("        <SYSTEM>", elements + "\n        <SYSTEM>", 1)
    if system_children:
        text = text.replace("<CATEGORY>SYSTEM_DESCRIPTION</CATEGORY>",
                            "<CATEGORY>SYSTEM_DESCRIPTION</CATEGORY>" + system_children, 1)
    if composition_children:
        marker = "<SHORT-NAME>Top</SHORT-NAME>"
        text = text.replace(marker, marker + composition_children, 1)
    path.write_text(text, encoding="utf-8")


def run_white(tmp_path, rows: dict, general: dict | None = None, **config):
    path = write_white_table(tmp_path, rows)
    if general:
        from openpyxl import load_workbook

        workbook = load_workbook(path)
        sheet = workbook["General"]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value in general:
                row[1].value = general[row[0].value]
        workbook.save(path)
        workbook.close()
    return run_tool(SYSTEM, "--config", write_config(tmp_path, **config))


def output_root(tmp_path) -> ET.Element:
    return ET.parse(tmp_path / "out" / "TSystem.arxml").getroot()


def tags(root: ET.Element, tag: str) -> list[ET.Element]:
    return list(root.iter(f"{{{AR}}}{tag}"))


def names(root: ET.Element, tag: str) -> list[str]:
    return [item.findtext("ar:SHORT-NAME", namespaces=NS) for item in tags(root, tag)]


def test_reference_only_system_has_no_empty_swc_to_ecu_mapping(tmp_path):
    edit_model(tmp_path)
    result = run_white(tmp_path, {"Components": REFERENCES, "AssemblySwConnection": [["A", "pOut", "B", "rIn"]]},
                       generateSystem=True, interfaceMismatch="skip")
    assert result.returncode == 0, result.stderr
    root = output_root(tmp_path)
    assert tags(root, "SWC-TO-ECU-MAPPING") == []
    assert tags(root, "SYSTEM-MAPPING") == []
    assert names(root, "ROOT-SW-COMPOSITION-PROTOTYPE") == ["RootProto"]


def test_added_component_without_implementation_only_warns(tmp_path):
    edit_model(tmp_path)
    result = run_white(tmp_path, {"Components": [["A2", "A2", "/Pkg/A", "Add"]]}, generateSystem=True)
    assert result.returncode == 0, result.stderr
    assert "has no SWC-INTERNAL-BEHAVIOR; SwcToImplMapping is not generated" in result.stderr
    root = output_root(tmp_path)
    assert names(root, "SW-COMPONENT-PROTOTYPE") == ["A2"]
    assert tags(root, "SWC-TO-IMPL-MAPPING") == []
    mapping = tags(root, "SWC-TO-ECU-MAPPING")
    assert len(mapping) == 1
    assert len(tags(mapping[0], "COMPONENT-IREF")) == 1


@pytest.mark.parametrize("system, expected", [("Sys", "/Pkg/Sys"), ("/Pkg/Sys", "/Pkg/Sys"), ("Nope", None)])
def test_system_short_name_resolves_to_existing_system(tmp_path, system, expected):
    edit_model(tmp_path)
    result = run_white(tmp_path, {"Components": [["A2", "A2", "/Pkg/A", "Add"]]},
                       general={"System": system}, generateSystem=True)
    if expected is None:
        assert result.returncode == 1
        assert "System is not defined in the loaded ARXML: Nope" in result.stderr
        return
    assert result.returncode == 0, result.stderr
    root = output_root(tmp_path)
    packages = [package.findtext("ar:SHORT-NAME", namespaces=NS) for package in tags(root, "AR-PACKAGE")]
    assert packages == ["Pkg"]
    assert names(root, "SYSTEM") == ["Sys"]
    assert "has no SYSTEM-VERSION" in result.stderr
    assert tags(root, "SYSTEM-VERSION") == []


@pytest.mark.parametrize("existing, written, warning", [
    ("", True, False),
    ("<FIBEX-ELEMENTS><FIBEX-ELEMENT-REF-CONDITIONAL><FIBEX-ELEMENT-REF DEST=\"ECU-INSTANCE\">/Pkg/Ecu"
     "</FIBEX-ELEMENT-REF></FIBEX-ELEMENT-REF-CONDITIONAL></FIBEX-ELEMENTS>", False, False),
    ("<FIBEX-ELEMENTS><FIBEX-ELEMENT-REF-CONDITIONAL><FIBEX-ELEMENT-REF DEST=\"ECU-INSTANCE\">/Pkg/Other"
     "</FIBEX-ELEMENT-REF></FIBEX-ELEMENT-REF-CONDITIONAL></FIBEX-ELEMENTS>", False, True),
])
def test_fibex_elements_are_not_split(tmp_path, existing, written, warning):
    edit_model(tmp_path, "<ECU-INSTANCE><SHORT-NAME>Other</SHORT-NAME></ECU-INSTANCE>", system_children=existing)
    result = run_white(tmp_path, {"Components": [["A2", "A2", "/Pkg/A", "Add"]]}, generateSystem=True)
    assert result.returncode == 0, result.stderr
    assert bool(tags(output_root(tmp_path), "FIBEX-ELEMENTS")) == written
    assert ("FIBEX-ELEMENTS cannot be split" in result.stderr) == warning


@pytest.mark.parametrize("scenario, message", [
    ("missing-ecu", "ECU instance is not defined: /Pkg/Missing"),
    ("ecu-is-signal", "ECU instance is not defined: /Pkg/Sig"),
    ("foreign-root", "does not belong to System /Pkg/Sys"),
    ("existing-system-mapping", "SystemMapping /Pkg/Sys/SystemMapping_TSystemImpl already exists"),
])
def test_system_references_are_checked(tmp_path, scenario, message):
    children = ""
    if scenario == "existing-system-mapping":
        children = "<MAPPINGS><SYSTEM-MAPPING><SHORT-NAME>SystemMapping_TSystemImpl</SHORT-NAME></SYSTEM-MAPPING></MAPPINGS>"
    edit_model(tmp_path, "<SYSTEM-SIGNAL><SHORT-NAME>Sig</SHORT-NAME></SYSTEM-SIGNAL>", system_children=children)
    general = {
        "missing-ecu": {"EcuInstance": "/Pkg/Missing"},
        "ecu-is-signal": {"EcuInstance": "/Pkg/Sig"},
        "foreign-root": {"RootSwCompositionPrototype": "/Pkg/Other/RootProto"},
    }.get(scenario)
    result = run_white(tmp_path, {"Components": [["A2", "A2", "/Pkg/A", "Add"]]}, general=general,
                       generateSystem=True)
    assert result.returncode == 1
    assert message in result.stderr
    assert not (tmp_path / "out").exists()


def test_case_only_system_mapping_name_is_a_warning(tmp_path):
    edit_model(tmp_path, system_children="<MAPPINGS><SYSTEM-MAPPING><SHORT-NAME>systemmapping_tsystemimpl"
                                         "</SHORT-NAME></SYSTEM-MAPPING></MAPPINGS>")
    result = run_white(tmp_path, {"Components": [["A2", "A2", "/Pkg/A", "Add"]]}, generateSystem=True)
    assert result.returncode == 0, result.stderr
    assert "SystemMapping SystemMapping_TSystemImpl differs only in case" in result.stderr
    assert names(output_root(tmp_path), "SYSTEM-MAPPING") == ["SystemMapping_TSystemImpl"]


@pytest.mark.parametrize("rows, existing, message", [
    ({"Components": [*REFERENCES, ["pTop", "pTop", "/Pkg/A", "Add"]],
      "DelegationSwConnection": [["A", "pOut", "pTop"]]}, "", "new port pTop conflicts with new prototype pTop"),
    ({"Components": [["A", "A", "/Pkg/B", "Add"]]}, "", "prototype type mismatch: /Pkg/Top/A"),
])
def test_composition_namespace_rejects_exact_duplicates(tmp_path, rows, existing, message):
    edit_model(tmp_path, composition_children=existing)
    result = run_white(tmp_path, rows)
    assert result.returncode == 1
    assert message in result.stderr
    assert not (tmp_path / "out").exists()


@pytest.mark.parametrize("rows, existing, message, generated", [
    ({"Components": [["A2", "A2", "/Pkg/A", "Add"], ["a2", "a2", "/Pkg/A", "Add"]]}, "",
     "a2 differs only in case from A2", ["A2", "a2"]),
    ({"Components": [["a", "a", "/Pkg/A", "Add"]]}, "", "a differs only in case from A", ["a"]),
    ({"Components": REFERENCES, "DelegationSwConnection": [["A", "pOut", "AC_A_pOut_B_rIn"]]},
     "<CONNECTORS><ASSEMBLY-SW-CONNECTOR><SHORT-NAME>ac_A_pOut_B_rIn</SHORT-NAME></ASSEMBLY-SW-CONNECTOR>"
     "</CONNECTORS>", "AC_A_pOut_B_rIn differs only in case from ac_A_pOut_B_rIn", []),
])
def test_case_only_names_are_warnings(tmp_path, rows, existing, message, generated):
    edit_model(tmp_path, composition_children=existing)
    result = run_white(tmp_path, rows)
    assert result.returncode == 0, result.stderr
    assert message in result.stderr
    assert names(output_root(tmp_path), "SW-COMPONENT-PROTOTYPE") == generated


def test_outer_port_name_used_by_other_element_is_rejected(tmp_path):
    edit_model(tmp_path, composition_children="<CONNECTORS><ASSEMBLY-SW-CONNECTOR><SHORT-NAME>x</SHORT-NAME>"
                                              "</ASSEMBLY-SW-CONNECTOR></CONNECTORS>")
    result = run_white(tmp_path, {"Components": REFERENCES, "DelegationSwConnection": [["A", "pOut", "x"]]})
    assert result.returncode == 1
    assert "outer port name '/Pkg/Top/x' is already used by a ASSEMBLY-SW-CONNECTOR" in result.stderr


def _assembly(name: str) -> str:
    return f"""<CONNECTORS><ASSEMBLY-SW-CONNECTOR><SHORT-NAME>{name}</SHORT-NAME>
      <PROVIDER-IREF>
        <CONTEXT-COMPONENT-REF DEST="SW-COMPONENT-PROTOTYPE">/Pkg/Top/A</CONTEXT-COMPONENT-REF>
        <TARGET-P-PORT-REF DEST="P-PORT-PROTOTYPE">/Pkg/A/pOut</TARGET-P-PORT-REF>
      </PROVIDER-IREF>
      <REQUESTER-IREF>
        <CONTEXT-COMPONENT-REF DEST="SW-COMPONENT-PROTOTYPE">/Pkg/Top/B</CONTEXT-COMPONENT-REF>
        <TARGET-R-PORT-REF DEST="R-PORT-PROTOTYPE">/Pkg/B/rIn</TARGET-R-PORT-REF>
      </REQUESTER-IREF>
    </ASSEMBLY-SW-CONNECTOR></CONNECTORS>"""


@pytest.mark.parametrize("existing_name, warning", [("ac_A_pOut_B_rIn", False), ("Manual", True)])
def test_existing_connection_is_reused(tmp_path, existing_name, warning):
    edit_model(tmp_path, composition_children=_assembly(existing_name))
    path = tmp_path / "model.arxml"
    path.write_text(path.read_text(encoding="utf-8").replace(">/Pkg/IfB<", ">/Pkg/IfA<"), encoding="utf-8")
    result = run_white(tmp_path, {"Components": REFERENCES, "AssemblySwConnection": [["A", "pOut", "B", "rIn"]]})
    assert result.returncode == 0, result.stderr
    assert tags(output_root(tmp_path), "ASSEMBLY-SW-CONNECTOR") == []
    assert ("connects the same ports as existing connector Manual" in result.stderr) == warning


def test_wrong_interface_dest_is_corrected_with_warning(tmp_path):
    edit_model(tmp_path)
    path = tmp_path / "model.arxml"
    text = path.read_text(encoding="utf-8").replace(
        'DEST="SENDER-RECEIVER-INTERFACE">/Pkg/IfA<', 'DEST="CLIENT-SERVER-INTERFACE">/Pkg/IfA<', 1)
    path.write_text(text, encoding="utf-8")
    result = run_white(tmp_path, {"Components": REFERENCES, "DelegationSwConnection": [["A", "pOut", "pTop"]]})
    assert result.returncode == 0, result.stderr
    assert "declares interface DEST CLIENT-SERVER-INTERFACE" in result.stderr
    reference = tags(output_root(tmp_path), "PROVIDED-INTERFACE-TREF")[0]
    assert reference.attrib["DEST"] == "SENDER-RECEIVER-INTERFACE"


def test_pr_port_is_reported_as_unsupported(tmp_path):
    edit_model(tmp_path)
    path = tmp_path / "model.arxml"
    text = path.read_text(encoding="utf-8").replace(
        "</PORTS>\n        </APPLICATION-SW-COMPONENT-TYPE>",
        "<PR-PORT-PROTOTYPE><SHORT-NAME>prBoth</SHORT-NAME></PR-PORT-PROTOTYPE></PORTS>\n"
        "        </APPLICATION-SW-COMPONENT-TYPE>", 1)
    path.write_text(text, encoding="utf-8")
    result = run_white(tmp_path, {"Components": REFERENCES, "DelegationSwConnection": [["A", "prBoth", "pTop"]]})
    assert result.returncode == 1
    assert "PR-PORT-PROTOTYPE '/Pkg/A/prBoth' is not supported by this tool" in result.stderr


def test_duplicate_signal_rows_and_group_type_mapping_warn(tmp_path):
    edit_model(tmp_path, "<SYSTEM-SIGNAL><SHORT-NAME>Sig</SHORT-NAME></SYSTEM-SIGNAL>",
               "<SYSTEM-SIGNAL-GROUP><SHORT-NAME>Grp</SHORT-NAME></SYSTEM-SIGNAL-GROUP>")
    result = run_white(tmp_path, {
        "Components": [["A", "A", "/Pkg/A", "Reference"]],
        "SenderReceiverToSignal": [["/Pkg/Sig", "OUT", "A", "pOut", "a"], ["/Pkg/Sig", "out", "A", "pOut", "a"]],
        "SenderReceiverToSignalGroup": [["/Pkg/Grp", "A", "pWide", "c"]],
    }, generateSystem=True)
    assert result.returncode == 0, result.stderr
    assert "duplicate signal mapping /Pkg/IfA/a -> /Pkg/Sig skipped" in result.stderr
    assert "requires a TYPE-MAPPING" in result.stderr
    root = output_root(tmp_path)
    assert len(tags(root, "SENDER-RECEIVER-TO-SIGNAL-MAPPING")) == 1
    assert len(tags(root, "SENDER-RECEIVER-TO-SIGNAL-GROUP-MAPPING")) == 1


@pytest.mark.parametrize("release", ["R19-11", "4.3.0"])
def test_ar_release_is_not_matched_against_inputs(tmp_path, release):
    edit_model(tmp_path)
    result = run_white(tmp_path, {"Components": REFERENCES}, arRelease=release)
    assert result.returncode == 0, result.stderr
    assert "release" not in result.stderr.lower()


def test_existing_generated_mapping_set_cannot_be_split(tmp_path):
    edit_model(tmp_path, "<PORT-INTERFACE-MAPPING-SET><SHORT-NAME>Top_InterfaceMappings</SHORT-NAME>"
                         "</PORT-INTERFACE-MAPPING-SET>")
    result = run_white(tmp_path, {"Components": REFERENCES, "AssemblySwConnection": [["A", "pOut", "B", "rIn"]]},
                       interfaceMismatch="map")
    assert result.returncode == 1
    assert "PORT-INTERFACE-MAPPING-SET /Pkg/Top_InterfaceMappings already exists" in result.stderr


def test_green_bus_channel_is_only_a_hint(tmp_path):
    edit_model(tmp_path)
    path = tmp_path / "model.arxml"
    text = path.read_text(encoding="utf-8")
    text = text.replace("      </ELEMENTS>\n    </AR-PACKAGE>\n    <AR-PACKAGE>\n      <SHORT-NAME>Other</SHORT-NAME>",
                        """        <SWC-IMPLEMENTATION><SHORT-NAME>AImpl</SHORT-NAME>
          <BEHAVIOR-REF DEST="SWC-INTERNAL-BEHAVIOR">/Pkg/A/IB</BEHAVIOR-REF></SWC-IMPLEMENTATION>
      </ELEMENTS>
    </AR-PACKAGE>
    <AR-PACKAGE><SHORT-NAME>Ch1</SHORT-NAME><AR-PACKAGES><AR-PACKAGE><SHORT-NAME>Signal</SHORT-NAME>
      <ELEMENTS><SYSTEM-SIGNAL><SHORT-NAME>Sig</SHORT-NAME></SYSTEM-SIGNAL></ELEMENTS>
    </AR-PACKAGE></AR-PACKAGES></AR-PACKAGE>
    <AR-PACKAGE>
      <SHORT-NAME>Other</SHORT-NAME>""", 1)
    text = text.replace("</PORTS>\n        </APPLICATION-SW-COMPONENT-TYPE>",
                        "</PORTS>\n          <INTERNAL-BEHAVIORS><SWC-INTERNAL-BEHAVIOR><SHORT-NAME>IB</SHORT-NAME>"
                        "</SWC-INTERNAL-BEHAVIOR></INTERNAL-BEHAVIORS>\n        </APPLICATION-SW-COMPONENT-TYPE>", 1)
    path.write_text(text, encoding="utf-8")
    workbook = Workbook()
    workbook.active.title = "R-Port"
    workbook.create_sheet("P-Port")
    for sheet in workbook:
        sheet.append(["", "Port", "Module", "", "", "", "", "DataElement", "Component", "Port", "Remark"])
    workbook["P-Port"].append(["", "pOut", "CAN", "", "", "", "", "a", "Ch1", "Sig", ""])
    workbook.save(tmp_path / "A.xlsx")
    workbook.close()
    config = write_config(tmp_path, generateSystem=True, greenTables=["A.xlsx"],
                          componentAliases={"A": "/Pkg/Top/A"},
                          general={"Composition": "/Pkg/Top", "System": "/Pkg/Sys",
                                   "RootSwCompositionPrototype": "/Pkg/Sys/RootProto", "EcuInstance": "/Pkg/Ecu"})
    contents = json.loads(config.read_text(encoding="utf-8"))
    del contents["whiteTable"]
    config.write_text(json.dumps(contents), encoding="utf-8")
    result = run_tool(SYSTEM, "--config", config)
    assert result.returncode == 0, result.stderr
    assert "channel 'Ch1' is not a CAN physical channel" in result.stderr
    reference = tags(output_root(tmp_path), "SYSTEM-SIGNAL-REF")[0]
    assert reference.text == "/Ch1/Signal/Sig"
